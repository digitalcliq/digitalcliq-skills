#!/usr/bin/env python3
"""Shared post-flight validator for DigitalCLIQ deliverable skills.

Recreated 2026-08-28 after the original was lost in the 2026-08-16 vault
migration (old vault `.claude/skills/post_flight.py` deleted; skills moved to
the skills-plugin directory). Lives as a sibling of the skill folders so every
skill's `_validate_deliverable()` finds it at its first lookup candidate; a
copy is kept at `/Users/drewmoon/Desktop/DigitalCLIQ Brain HQ/.claude/skills/`
(the second candidate) for CLI use from SKILL.md instructions.

Contract (per score-leads SKILL.md "Self-Validation"):

    from post_flight import validate_report
    result = validate_report(output_path, min_sheets=4)
    result["passed"]   -> bool
    result["summary"]  -> str
    result["checks"]   -> list of {name, passed, detail}

Four checks, all must pass:
  1. branding      - a DigitalCLIQ logo image is actually embedded in the file
                     (xl/media for xlsx, word/media for docx, image XObjects
                     for pdf). A bare styled masthead with no image FAILS.
  2. location      - the file lives inside the vault outputs staging folder:
                     /Users/drewmoon/Desktop/DigitalCLIQ Brain HQ/outputs/
  3. min sections  - at least `min_sheets` sheets (xlsx), `min_pages` pages
                     (pdf), or heading sections (docx).
  4. no placeholders - no template tokens ({dealer_name}-style), literal
                     YYYY-MM-DD, PLACEHOLDER, or lorem ipsum in the content.

CLI (used by SKILL.md validation blocks):

    python3 post_flight.py "<output path>" [--min-sheets N] [--min-pages N]

Prints one line per check plus a summary line, exits 0 on pass, 2 on fail.
"""

import os
import re
import sys
import zipfile

OUTPUTS_DIR = "/Users/drewmoon/Desktop/DigitalCLIQ Brain HQ/outputs"

# Template tokens like {dealer_name} / {client}, plus the literal junk that
# means a data section never got filled in.
_PLACEHOLDER_RES = [
    (re.compile(r"\{[a-z][a-z0-9_]*\}"), "template token like {dealer_name}"),
    (re.compile(r"YYYY-MM-DD"), "literal YYYY-MM-DD"),
    (re.compile(r"placeholder", re.IGNORECASE), "PLACEHOLDER text"),
    (re.compile(r"lorem\s+ipsum", re.IGNORECASE), "lorem ipsum"),
]

_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".emf", ".wmf", ".tiff")


# ---------------------------------------------------------------- extractors

def _xlsx_facts(path, min_sheets, _min_pages):
    """(has_logo, section_count, section_label, text) for an Excel workbook."""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        has_logo = any(
            n.startswith("xl/media/") and n.lower().endswith(_IMAGE_EXTS)
            for n in names
        )

    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_count = len(wb.sheetnames)
        chunks = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        chunks.append(v)
    finally:
        wb.close()
    return has_logo, sheet_count, "sheets", "\n".join(chunks)


def _pdf_facts(path, _min_sheets, _min_pages):
    """(has_logo, page_count, 'pages', text) for a PDF via PyMuPDF."""
    import fitz  # PyMuPDF
    doc = fitz.open(path)
    try:
        page_count = doc.page_count
        has_logo = any(doc.get_page_images(i) for i in range(page_count))
        text = "\n".join(doc.load_page(i).get_text() for i in range(page_count))
    finally:
        doc.close()
    return has_logo, page_count, "pages", text


def _docx_facts(path, _min_sheets, _min_pages):
    """(has_logo, heading_section_count, 'sections', text) for a Word doc."""
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        has_logo = any(
            n.startswith("word/media/") and n.lower().endswith(_IMAGE_EXTS)
            for n in names
        )

    import docx
    d = docx.Document(path)
    chunks, sections = [], 0
    for p in d.paragraphs:
        if p.text:
            chunks.append(p.text)
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading") and p.text.strip():
            sections += 1
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    chunks.append(cell.text)
    return has_logo, sections, "heading sections", "\n".join(chunks)


_EXTRACTORS = {
    ".xlsx": _xlsx_facts,
    ".xlsm": _xlsx_facts,
    ".pdf": _pdf_facts,
    ".docx": _docx_facts,
}


# ------------------------------------------------------------------- checks

def validate_report(output_path, min_sheets=4, min_pages=None):
    """Validate a finished deliverable. Returns a dict with at least
    `passed` (bool) and `summary` (str); `checks` carries per-check detail."""
    checks = []

    output_path = os.path.abspath(output_path)
    if not os.path.isfile(output_path):
        summary = f"post-flight FAILED: file not found: {output_path}"
        return {"passed": False, "summary": summary, "checks": [
            {"name": "exists", "passed": False, "detail": summary}]}

    ext = os.path.splitext(output_path)[1].lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is None:
        summary = (f"post-flight FAILED: unsupported file type '{ext}' "
                   f"(supported: {', '.join(sorted(_EXTRACTORS))})")
        return {"passed": False, "summary": summary, "checks": [
            {"name": "filetype", "passed": False, "detail": summary}]}

    try:
        has_logo, section_count, section_label, text = extractor(
            output_path, min_sheets, min_pages)
    except Exception as e:  # corrupt/unreadable file is a hard fail, not a skip
        summary = f"post-flight FAILED: could not read {output_path}: {e}"
        return {"passed": False, "summary": summary, "checks": [
            {"name": "readable", "passed": False, "detail": summary}]}

    # 1. branding: logo image actually embedded
    checks.append({
        "name": "branding",
        "passed": has_logo,
        "detail": ("logo image embedded" if has_logo else
                   "no embedded image found (DigitalCLIQ logo missing; a bare "
                   "styled masthead with no logo image is a failure, not a "
                   "fallback; canonical logo: "
                   f"{OUTPUTS_DIR.rsplit('/outputs', 1)[0]}"
                   "/Resources/brand-assets/digital-cliq-logo-solid-1000px-wide.png)"),
    })

    # 2. location: inside the vault outputs staging folder
    real = os.path.realpath(output_path)
    in_outputs = real.startswith(os.path.realpath(OUTPUTS_DIR) + os.sep)
    checks.append({
        "name": "location",
        "passed": in_outputs,
        "detail": (f"inside {OUTPUTS_DIR}/" if in_outputs else
                   f"file is at {real}, not inside {OUTPUTS_DIR}/ "
                   "(move it into outputs/ staging and re-validate)"),
    })

    # 3. minimum sheet / page / section count
    min_required = min_pages if (ext == ".pdf" and min_pages) else min_sheets
    enough = section_count >= min_required
    checks.append({
        "name": "min_sections",
        "passed": enough,
        "detail": (f"{section_count} {section_label} (>= {min_required} required)"
                   if enough else
                   f"only {section_count} {section_label}, {min_required} required "
                   "(a section likely failed silently or the report is truncated)"),
    })

    # 4. no placeholder tokens
    hits = []
    for rx, label in _PLACEHOLDER_RES:
        m = rx.search(text)
        if m:
            hits.append(f"{label} ('{m.group(0)}')")
    checks.append({
        "name": "no_placeholders",
        "passed": not hits,
        "detail": ("no placeholder tokens" if not hits else
                   "placeholder content found: " + "; ".join(hits) +
                   " (a data section came back empty; re-fetch the data, "
                   "do not paper over it)"),
    })

    passed = all(c["passed"] for c in checks)
    if passed:
        summary = (f"post-flight PASSED: logo embedded, in outputs/, "
                   f"{section_count} {section_label}, no placeholders")
    else:
        fails = [f"{c['name']}: {c['detail']}" for c in checks if not c["passed"]]
        summary = "post-flight FAILED: " + " | ".join(fails)

    return {"passed": passed, "summary": summary, "checks": checks}


# ---------------------------------------------------------------------- CLI

def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    path = argv[0]
    min_sheets, min_pages = 4, None
    args = argv[1:]
    i = 0
    while i < len(args):
        if args[i] == "--min-sheets" and i + 1 < len(args):
            min_sheets = int(args[i + 1]); i += 2
        elif args[i] == "--min-pages" and i + 1 < len(args):
            min_pages = int(args[i + 1]); i += 2
        else:
            print(f"ERROR: unknown argument {args[i]!r}", file=sys.stderr)
            return 1

    result = validate_report(path, min_sheets=min_sheets, min_pages=min_pages)
    for c in result.get("checks", []):
        mark = "PASS" if c["passed"] else "FAIL"
        print(f"  [{mark}] {c['name']}: {c['detail']}")
    print(result["summary"])
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
