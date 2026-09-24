"""DigitalCLIQ Social Media Manager — branded Excel workbook builder.

Renders a single plan JSON (see CONTRACT below) into a DigitalCLIQ-branded
.xlsx with embedded asset thumbnails, clickable links, and a contact tab.

Reusable for ANY client. The /social-media-manager skill fills the JSON; this
script is pure rendering and does no strategy work.

Usage:
    python3 build_workbook.py <plan.json> <output.xlsx>

The plan JSON contract (every key optional except meta + scoreboard):

{
  "meta": {
    "client": "McPeek Chrysler Dodge RAM of Anaheim",
    "business_type": "Automotive — CDJR / RAM dealership",
    "period_label": "90-Day Growth Plan • Jun–Sep 2026",
    "generated_on": "2026-06-24",
    "goal": "2x followers and interactions across IG, TikTok, YouTube, X/Threads, Facebook in 90 days",
    "prepared_by": "DigitalCLIQ — Digital Strategy & Development"
  },
  "scoreboard": [
    {"platform","handle","baseline_followers","baseline_engagement_rate",
     "baseline_avg_interactions","target_30_followers","target_60_followers",
     "target_90_followers","twox_followers","twox_interactions","primary_levers"}
  ],
  "playbook": [
    {"platform","ranking_signals":[...],"winning_formats":[...],"cadence",
     "what_changed_2026":[...],"beyond_hashtags","automotive_angle":[...]}
  ],
  "calendar": [
    {"week","phase","date","platform","campaign","format","hook",
     "caption_direction","gen_tool","prompt","kpi","status",
     "asset_thumb": "/abs/path/thumb.png", "asset_link": "/abs/path/full.mp4"}
  ],
  "campaigns": [
    {"name","phase","platforms","big_idea","why_it_works","execution","kpis","assets_needed"}
  ],
  "prompt_library": [
    {"asset_id","title","type","platform","magnific_model","settings","prompt",
     "fallback_claude","fallback_gemini","fallback_sora"}
  ],
  "generated_assets": [
    {"asset_id","title","type","magnific_model","credits","vault_path",
     "web_url","thumb_path","prompt"}
  ],
  "contact": {
    "company","tagline","contact_name","title","email","phone","website","note"
  }
}
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

for _pkg in ("openpyxl", "PIL"):
    try:
        __import__(_pkg)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install",
                               "openpyxl" if _pkg == "openpyxl" else "Pillow"])

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

# ── DigitalCLIQ design-system tokens (Resources/design-system/Design-System.md) ──
DIGITAL_BLUE = "405FAB"
SKY_BLUE = "6B9DD4"
NAVY_DEEP = "070A15"
INK = NAVY_DEEP           # body text on light backgrounds
WARM_GREY = "949592"
CALLOUT_TINT = "EDF2F9"   # alternating body rows / light fills
BORDER_BLUE = "D8E1F0"    # thin borders + deeper light tint
TILE_BLUE = "2E4780"

# Palette tints only. The phase/status TEXT in the cell carries the meaning;
# fills are treatment (Sky Blue family = strong, Warm Grey = weak, Digital Blue = emphasis).
PHASE_FILL = {"0-30": CALLOUT_TINT, "30-60": BORDER_BLUE, "60-90": SKY_BLUE}
STATUS_FILL = {
    "generated": SKY_BLUE, "ready": SKY_BLUE, "prompt only": WARM_GREY,
    "draft": WARM_GREY, "scheduled": DIGITAL_BLUE, "queued": WARM_GREY,
}

LOGO_PATH = Path("/Users/drewmoon/Desktop/DigitalCLIQ Brain HQ/Resources/brand-assets/"
                 "digital-cliq-logo-solid-1000px-wide.png")

# Hard guardrail: this address must NEVER appear in any DigitalCLIQ deliverable.
BANNED_EMAIL = "digitalcliq@gmail.com"
APPROVED_EMAIL = "drewmoon@digitalcliq.com"


# ── style helpers ─────────────────────────────────────────────────────
# Design-System fonts: Dosis = structure (headings, stats, labels),
# Roboto Slab = reading (body). Arial/Calibri/Helvetica in output = FAIL.
def _f_header():  return Font(name="Dosis", size=10, bold=True, color="FFFFFF")
def _f_body(b=False): return Font(name="Roboto Slab", size=10, bold=b, color=INK)
def _f_muted():   return Font(name="Roboto Slab", size=9, italic=True, color=WARM_GREY)
def _f_title():   return Font(name="Dosis", size=20, bold=True, color="FFFFFF")
def _f_subtitle(): return Font(name="Dosis", size=12, bold=True, color=SKY_BLUE)
def _f_section(): return Font(name="Dosis", size=13, bold=True, color=DIGITAL_BLUE)
def _f_link():    return Font(name="Roboto Slab", size=10, underline="single", color=DIGITAL_BLUE)
def _fill(c):     return PatternFill("solid", fgColor=c)
def _border():
    s = Side(style="thin", color=BORDER_BLUE)
    return Border(left=s, right=s, top=s, bottom=s)


def _logo_image(width: int = 104) -> XLImage:
    """White knockout logo from the canonical path, at true 500x154 aspect.
    A missing logo fails loudly, never silently."""
    if not LOGO_PATH.exists():
        raise FileNotFoundError(
            f"DigitalCLIQ logo not found at canonical path: {LOGO_PATH}. "
            "Fix the path; do not ship without branding.")
    img = XLImage(str(LOGO_PATH))
    img.width, img.height = width, int(width * 154 / 500)
    return img


def _masthead(ws, title, ncols, date="", title_col=3):
    """Design-System §4 Excel masthead: rows 1-2 Digital Blue band, white logo
    anchored A1 (~0.35in tall), sheet title in white Dosis 14 Bold, date under it."""
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    for r in (1, 2):
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = _fill(DIGITAL_BLUE)
    ws.add_image(_logo_image(104), "A1")
    t = ws.cell(row=1, column=title_col, value=title)
    t.font = Font(name="Dosis", size=14, bold=True, color="FFFFFF")
    t.alignment = Alignment(vertical="center")
    if date:
        d = ws.cell(row=2, column=title_col, value=date)
        d.font = Font(name="Dosis", size=9, color=BORDER_BLUE)
        d.alignment = Alignment(vertical="center")


def _hrow(ws, headers, row=1, fill=DIGITAL_BLUE, height=30):
    ws.row_dimensions[row].height = height
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _f_header()
        c.fill = _fill(fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border()


def _brow(ws, row_idx, values, wrap_over=28):
    fill = _fill(CALLOUT_TINT) if row_idx % 2 == 0 else None
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=col, value=v)
        c.font = _f_body()
        c.alignment = Alignment(vertical="top",
                                wrap_text=isinstance(v, str) and len(str(v)) > wrap_over)
        c.border = _border()
        if fill:
            c.fill = fill


def _bullets(items):
    if not items:
        return ""
    if isinstance(items, str):
        return items
    return "\n".join(f"• {x}" for x in items)


def _thumb_for(path: str, max_px: int = 150) -> str | None:
    """Return a path to a <=max_px PNG thumbnail for embedding. Builds one if needed."""
    if not path or not os.path.exists(path):
        return None
    ext = os.path.splitext(path)[1].lower()
    # Embeddable still images only. Videos must be passed as a poster image path.
    if ext not in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return None
    try:
        im = PILImage.open(path).convert("RGB")
        im.thumbnail((max_px, max_px))
        tdir = Path(path).parent / "_xlsx_thumbs"
        tdir.mkdir(exist_ok=True)
        out = tdir / (Path(path).stem + f"_{max_px}.png")
        im.save(out, "PNG")
        return str(out)
    except Exception:
        return None


# ── Tab 1: Cover ──────────────────────────────────────────────────────
def _cover(wb, meta):
    ws = wb.create_sheet("Cover")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 96

    for r in range(1, 11):
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = _fill(NAVY_DEEP)
        ws.row_dimensions[r].height = 24

    ws.add_image(_logo_image(210), "B2")

    ws["B5"] = "Social Media Growth Plan"
    ws["B5"].font = _f_title()
    ws["B6"] = meta.get("client", "")
    ws["B6"].font = Font(name="Dosis", size=15, bold=True, color="FFFFFF")
    ws["B7"] = meta.get("business_type", "")
    ws["B7"].font = Font(name="Dosis", size=11, color=SKY_BLUE)
    ws["B8"] = meta.get("period_label", "")
    ws["B8"].font = Font(name="Dosis", size=11, color="FFFFFF")

    ws["B12"] = "The Goal"
    ws["B12"].font = _f_section()
    ws["B13"] = meta.get("goal", "")
    ws["B13"].font = _f_body()
    ws["B13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[13].height = 44

    notes = [
        ("How to read this workbook", ""),
        ("Growth Scoreboard", "Each account's live baseline, the 30/60/90-day targets, and the 2x line. This is the number we are driving."),
        ("2026 Playbook", "Per-platform ranking signals and formats that actually move reach THIS year. Built without leaning on hashtags."),
        ("90-Day Calendar", "Every planned post: campaign, format, hook, the generation prompt, and a thumbnail/link to the asset we generated for it."),
        ("Campaign Deep-Dives", "The thinking behind each major campaign — why it works and how it compounds toward the 2x."),
        ("Magnific Prompt Library", "Paste-ready prompts with the exact Magnific model + settings, plus Claude / Gemini / Sora fallbacks."),
        ("Generated Assets", "The real photos and videos generated via Magnific, saved into the vault and linked here."),
        ("Contact", "Your DigitalCLIQ point of contact."),
    ]
    r = 15
    for label, val in notes:
        ws[f"B{r}"] = label if not val else f"{label}  —  {val}"
        ws[f"B{r}"].font = _f_body(b=True) if not val else _f_body()
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        if val:
            ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws[f"B{r}"] = f"Prepared by {meta.get('prepared_by', 'DigitalCLIQ — Digital Strategy & Development')}"
    ws[f"B{r}"].font = Font(name="Dosis", size=10, bold=True, color=DIGITAL_BLUE)
    ws[f"B{r+1}"] = f"Generated {meta.get('generated_on', '')}  •  {APPROVED_EMAIL}"
    ws[f"B{r+1}"].font = _f_muted()


# ── Tab 2: Growth Scoreboard (baseline → 2x) ──────────────────────────
def _scoreboard(wb, rows, meta):
    ws = wb.create_sheet("Growth Scoreboard")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Growth Scoreboard — Baseline to 2x", ncols=10,
              date=meta.get("generated_on", ""))
    ws["A3"] = "Baselines pulled live. Targets are the path to doubling followers and interactions in 90 days."
    ws["A3"].font = _f_muted()
    ws.merge_cells("A3:J3")

    headers = ["Platform", "Handle", "Baseline\nFollowers", "Baseline\nEng. Rate",
               "Baseline Avg\nInteractions/Post", "Day 30\nTarget", "Day 60\nTarget",
               "Day 90\nTarget (2x)", "2x\nInteractions", "Primary Growth Levers"]
    _hrow(ws, headers, row=4)

    r = 5
    for row in rows:
        vals = [
            row.get("platform", ""), row.get("handle", ""),
            row.get("baseline_followers", ""), row.get("baseline_engagement_rate", ""),
            row.get("baseline_avg_interactions", ""), row.get("target_30_followers", ""),
            row.get("target_60_followers", ""),
            row.get("target_90_followers", row.get("twox_followers", "")),
            row.get("twox_interactions", ""), row.get("primary_levers", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.cell(row=r, column=8).font = Font(name="Dosis", size=10, bold=True, color=DIGITAL_BLUE)
        r += 1

    widths = [13, 18, 11, 10, 13, 10, 10, 12, 12, 52]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"

    note = ("2x thesis: doubling is driven by reach per post (algorithm signals), not by posting more of the same. "
            "Each platform's levers above are sequenced across the 90-day calendar.")
    ws.cell(row=r + 1, column=1, value=note).font = _f_muted()
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=10)


# ── Tab 3: 2026 Playbook ──────────────────────────────────────────────
def _playbook(wb, rows):
    ws = wb.create_sheet("2026 Playbook")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "2026 Platform Playbook", ncols=6)
    ws["A3"] = "What ranks and what wins on each platform this year. Hashtags are not the strategy here."
    ws["A3"].font = _f_muted()
    ws.merge_cells("A3:F3")

    headers = ["Platform", "What The Algorithm Rewards (2026)", "Winning Formats",
               "Cadence", "What Changed in 2026", "Beyond Hashtags / Discovery"]
    _hrow(ws, headers, row=4)

    r = 5
    for row in rows:
        vals = [
            row.get("platform", ""),
            _bullets(row.get("ranking_signals")),
            _bullets(row.get("winning_formats")),
            row.get("cadence", ""),
            _bullets(row.get("what_changed_2026")),
            row.get("beyond_hashtags", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.row_dimensions[r].height = 150
        r += 1

    widths = [13, 42, 34, 22, 38, 40]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"


# ── Tab 4: 90-Day Calendar (with thumbnails + links) ──────────────────
def _calendar(wb, rows):
    ws = wb.create_sheet("90-Day Calendar")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "90-Day Calendar", ncols=14, title_col=4)
    headers = ["Wk", "Phase", "Target Date", "Platform", "Campaign", "Format",
               "Hook / Concept", "Caption / Script Direction", "Gen Tool + Model",
               "Generation Prompt", "KPI", "Asset", "Link", "Status"]
    _hrow(ws, headers, row=3)

    r = 4
    for row in rows:
        status = str(row.get("status", "")).strip().lower()
        vals = [
            row.get("week", ""), row.get("phase", ""), row.get("date", ""),
            row.get("platform", ""), row.get("campaign", ""), row.get("format", ""),
            row.get("hook", ""), row.get("caption_direction", ""),
            row.get("gen_tool", ""), row.get("prompt", ""), row.get("kpi", ""),
            "",  # asset thumbnail col (12)
            "",  # link col (13)
            row.get("status", ""),
        ]
        _brow(ws, r, vals)
        # phase tint
        pc = PHASE_FILL.get(str(row.get("phase", "")).strip())
        if pc:
            ws.cell(row=r, column=2).fill = _fill(pc)
        # status tint
        sc = STATUS_FILL.get(status)
        if sc:
            sc_cell = ws.cell(row=r, column=14)
            sc_cell.fill = _fill(sc)
            sc_cell.font = Font(name="Dosis", size=9, bold=True, color="FFFFFF")
            sc_cell.alignment = Alignment(horizontal="center", vertical="center")

        # embedded thumbnail
        thumb = _thumb_for(row.get("asset_thumb", ""))
        link = row.get("asset_link", "")
        if thumb:
            ws.row_dimensions[r].height = 120
            try:
                xi = XLImage(thumb)
                # scale to fit ~150px column / 120px row
                w, h = PILImage.open(thumb).size
                scale = min(150 / w, 110 / h)
                xi.width, xi.height = int(w * scale), int(h * scale)
                ws.add_image(xi, f"L{r}")
            except Exception:
                ws.cell(row=r, column=12, value="(asset)").font = _f_muted()
        else:
            ws.row_dimensions[r].height = 90
        if link:
            lc = ws.cell(row=r, column=13, value="Open asset")
            lc.hyperlink = link
            lc.font = _f_link()
            lc.alignment = Alignment(horizontal="center", vertical="center")
        r += 1

    widths = [5, 7, 12, 11, 22, 12, 30, 38, 18, 50, 20, 22, 12, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "C4"


# ── Tab 5: Campaign Deep-Dives ────────────────────────────────────────
def _campaigns(wb, rows):
    ws = wb.create_sheet("Campaign Deep-Dives")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Campaign Deep-Dives", ncols=8)
    headers = ["Campaign", "Phase", "Platforms", "The Big Idea", "Why It Works (2026)",
               "Execution", "KPIs", "Assets Needed"]
    _hrow(ws, headers, row=3)

    r = 4
    for row in rows:
        vals = [
            row.get("name", ""), row.get("phase", ""), row.get("platforms", ""),
            row.get("big_idea", ""), row.get("why_it_works", ""),
            row.get("execution", ""), row.get("kpis", ""), row.get("assets_needed", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.row_dimensions[r].height = 150
        r += 1

    widths = [24, 8, 16, 40, 40, 44, 26, 30]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"


# ── Tab 6: Magnific Prompt Library ────────────────────────────────────
def _prompt_library(wb, rows):
    ws = wb.create_sheet("Magnific Prompt Library")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Magnific Prompt Library — paste-ready", ncols=8)
    ws["A3"] = ("Primary path is Magnific (model + settings given). Fallback columns work in Claude, "
                "Gemini / Nano Banana, or Sora / Veo if generating outside Magnific.")
    ws["A3"].font = _f_muted()
    ws.merge_cells("A3:H3")

    headers = ["ID", "Asset Title", "Type", "Platform", "Magnific Model",
               "Settings", "Primary Prompt (Magnific)", "Fallback (Claude / Gemini / Sora)"]
    _hrow(ws, headers, row=4)

    r = 5
    for row in rows:
        fb = []
        if row.get("fallback_claude"): fb.append(f"Claude/Gemini image: {row['fallback_claude']}")
        if row.get("fallback_gemini"): fb.append(f"Gemini/Nano Banana: {row['fallback_gemini']}")
        if row.get("fallback_sora"): fb.append(f"Sora/Veo video: {row['fallback_sora']}")
        vals = [
            row.get("asset_id", ""), row.get("title", ""), row.get("type", ""),
            row.get("platform", ""), row.get("magnific_model", ""),
            row.get("settings", ""), row.get("prompt", ""), "\n\n".join(fb),
        ]
        _brow(ws, r, vals)
        ws.row_dimensions[r].height = 130
        r += 1

    widths = [6, 26, 10, 12, 22, 18, 56, 50]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"


# ── Tab 7: Generated Assets gallery ───────────────────────────────────
def _generated(wb, rows):
    ws = wb.create_sheet("Generated Assets")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Generated Assets — live from Magnific, saved to the vault", ncols=8, title_col=2)

    headers = ["Preview", "ID", "Title", "Type", "Magnific Model",
               "Credits", "Vault Path", "Web Link"]
    _hrow(ws, headers, row=3)

    r = 4
    for row in rows:
        vals = [
            "", row.get("asset_id", ""), row.get("title", ""), row.get("type", ""),
            row.get("magnific_model", ""), row.get("credits", ""),
            row.get("vault_path", ""), "",
        ]
        _brow(ws, r, vals)
        thumb = _thumb_for(row.get("thumb_path", ""), max_px=200)
        if thumb:
            ws.row_dimensions[r].height = 150
            try:
                xi = XLImage(thumb)
                w, h = PILImage.open(thumb).size
                scale = min(190 / w, 140 / h)
                xi.width, xi.height = int(w * scale), int(h * scale)
                ws.add_image(xi, f"A{r}")
            except Exception:
                pass
        else:
            ws.row_dimensions[r].height = 30
        if row.get("web_url"):
            lc = ws.cell(row=r, column=8, value="Open")
            lc.hyperlink = row["web_url"]
            lc.font = _f_link()
        r += 1

    widths = [28, 6, 26, 10, 22, 9, 48, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"


# ── Tab 8: Contact ────────────────────────────────────────────────────
def _contact(wb, c):
    ws = wb.create_sheet("Contact")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 56

    for r in range(1, 8):
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = _fill(NAVY_DEEP)
        ws.row_dimensions[r].height = 22
    ws.add_image(_logo_image(180), "B2")

    ws["B5"] = c.get("company", "DigitalCLIQ")
    ws["B5"].font = Font(name="Dosis", size=16, bold=True, color="FFFFFF")
    ws["B6"] = c.get("tagline", "Digital Strategy & Development")
    ws["B6"].font = Font(name="Dosis", size=11, color=SKY_BLUE)

    email = c.get("email", APPROVED_EMAIL)
    rows = [
        ("Your Contact", c.get("contact_name", "Drew Moon")),
        ("Title", c.get("title", "President")),
        ("Email", email),
        ("Phone", c.get("phone", "")),
        ("Website", c.get("website", "digitalcliq.com")),
    ]
    r = 9
    for label, val in rows:
        if not val:
            continue
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = _f_body(b=True)
        cell = ws[f"C{r}"]
        cell.value = val
        cell.font = _f_body()
        if label == "Email":
            cell.hyperlink = f"mailto:{val}"
            cell.font = _f_link()
        if label == "Website":
            cell.hyperlink = f"https://{val.lstrip('https://').lstrip('http://')}"
            cell.font = _f_link()
        r += 1

    if c.get("note"):
        r += 1
        ws[f"B{r}"] = c["note"]
        ws[f"B{r}"].font = _f_muted()
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)


# ── Public ────────────────────────────────────────────────────────────
_PLAYBOOK_DATA = Path(__file__).resolve().parent.parent / "references" / "playbook-data.json"


def _auto_fallbacks(a):
    t = a.get("type", "Image")
    return {
        "fallback_claude": a.get("fallback_claude",
            "Claude / Nano Banana image: same brand system + subject; strictly on-brand, no competitor brands or logos."),
        "fallback_gemini": a.get("fallback_gemini",
            "Gemini / Nano Banana: same prompt + 'photorealistic, on-brand lighting'; use a brand-safe reference for fidelity."),
        "fallback_sora": a.get("fallback_sora",
            ("Sora / Veo: animate the still — " + a.get("title", "")) if t in ("Image", "Graphic")
            else "Sora / Veo: same shot brief, 9:16, ~5s, subtle ambient audio."),
    }


def _expand_plan(plan: dict) -> dict:
    """Auto-derive prompt_library, generated_assets, calendar asset wiring, and the
    default 2026 playbook from compact inputs, so each run authors minimal data.

    Compact inputs the skill provides:
      plan['assets'] = [{id, title, type, platform, model, settings, prompt,
                         file, thumb, web, credits}]  (file/thumb filled by download_assets.py)
      calendar rows reference an asset by {"asset_id": "<id>"}
      plan['playbook_platforms'] = ["Instagram","TikTok",...] -> rows from playbook-data.json
      plan['playbook_overrides'] = {"Instagram": {...}}  (optional per-platform tweaks)
    """
    plan = dict(plan)
    assets = plan.get("assets")
    if assets:
        amap = {a["id"]: a for a in assets}
        if not plan.get("prompt_library"):
            plan["prompt_library"] = [dict(
                asset_id=a["id"], title=a.get("title", ""), type=a.get("type", ""),
                platform=a.get("platform", ""), magnific_model=a.get("model", ""),
                settings=a.get("settings", ""), prompt=a.get("prompt", ""), **_auto_fallbacks(a))
                for a in assets]
        if not plan.get("generated_assets"):
            plan["generated_assets"] = [dict(
                asset_id=a["id"], title=a.get("title", ""), type=a.get("type", ""),
                magnific_model=a.get("model", ""), credits=a.get("credits", ""),
                vault_path=a.get("file", ""), web_url=a.get("web", ""),
                thumb_path=a.get("thumb", ""), prompt=a.get("prompt", ""))
                for a in assets]
        for row in plan.get("calendar", []):
            aid = row.get("asset_id")
            if aid and aid in amap and not row.get("asset_thumb"):
                row["asset_thumb"] = amap[aid].get("thumb", "")
                row["asset_link"] = amap[aid].get("file", "")
    if not plan.get("playbook") and plan.get("playbook_platforms"):
        try:
            data = json.loads(_PLAYBOOK_DATA.read_text(encoding="utf-8"))
        except Exception:
            data = {}
        overrides = plan.get("playbook_overrides", {})
        pb = []
        for key in plan["playbook_platforms"]:
            entry = dict(data.get(key, {}))
            if not entry:
                continue
            entry["platform"] = entry.get("platform", key)
            entry.update(overrides.get(key, {}))
            pb.append(entry)
        if pb:
            plan["playbook"] = pb
    return plan


def _validate(path: str) -> list:
    """Post-flight: in outputs/, branded, no banned email, no placeholder tokens."""
    import zipfile
    issues = []
    if "/outputs/" not in os.path.abspath(path).replace("\\", "/"):
        issues.append("file is not inside outputs/")
    wb2 = load_workbook(path)
    imgs = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("xl/media")]
    if len(imgs) < 2:
        issues.append(f"only {len(imgs)} embedded images (logo or thumbnails missing)")
    blob = ""
    for ws in wb2.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str):
                    blob += c.value.lower() + " "
    if BANNED_EMAIL in blob:
        issues.append(f"BANNED {BANNED_EMAIL} present")
    # Real unfilled-template markers only (the descriptive word "placeholder" is fine in prose).
    for tok in ("lorem ipsum", "yyyy-mm-dd", "{client}", "{dealer", "tbd_", "xxxx@"):
        if tok in blob:
            issues.append(f"unfilled template token '{tok}'")
    return issues


def build(plan: dict, output_path: str) -> str:
    plan = _expand_plan(plan)
    # Guardrail: refuse to ship the banned address anywhere in the workbook.
    blob = json.dumps(plan).lower()
    if BANNED_EMAIL in blob:
        raise ValueError(
            f"REFUSED: '{BANNED_EMAIL}' appears in the plan data. DigitalCLIQ "
            f"deliverables use only {APPROVED_EMAIL}. Scrub the data and rebuild.")

    wb = Workbook()
    wb.remove(wb.active)
    _cover(wb, plan.get("meta", {}))
    _scoreboard(wb, plan.get("scoreboard", []), plan.get("meta", {}))
    if plan.get("playbook"):
        _playbook(wb, plan["playbook"])
    if plan.get("calendar"):
        _calendar(wb, plan["calendar"])
    if plan.get("campaigns"):
        _campaigns(wb, plan["campaigns"])
    if plan.get("prompt_library"):
        _prompt_library(wb, plan["prompt_library"])
    if plan.get("generated_assets"):
        _generated(wb, plan["generated_assets"])
    _contact(wb, plan.get("contact", {}))

    # Kill openpyxl's Calibri default on every touched cell
    # (design system: Arial/Calibri/Helvetica in output = FAIL).
    brand_default = Font(name="Roboto Slab", size=10, color=INK)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name in (None, "Calibri"):
                    cell.font = brand_default

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 build_workbook.py <plan.json> <output.xlsx>")
        sys.exit(2)
    plan_path, out_path = sys.argv[1], sys.argv[2]
    with open(plan_path, encoding="utf-8") as fh:
        plan_data = json.load(fh)
    saved = build(plan_data, out_path)
    print(f"Workbook written: {saved}")
    problems = _validate(saved)
    if problems:
        print("VALIDATION ISSUES:")
        for pr in problems:
            print("  -", pr)
        sys.exit(1)
    print("Validation OK: in outputs/, branded, no banned email, no placeholders.")
