"""DigitalCLIQ Blog Content Engine — branded Excel tracker builder.

Renders one plan JSON (see CONTRACT below) into a DigitalCLIQ-branded .xlsx that
acts as the index / control surface for a batch of blog posts and landing pages:
the content plan, the SEO + keyword brief, the AI-search (AEO/GEO) brief, the
article outlines, the Magnific prompt library, and the generated-asset gallery.

The full written body of each piece does NOT live here — it ships as a polished
branded .docx per piece (see write_docs.py). This workbook is the tracker + the
strategy intelligence behind every piece.

Reusable for ANY client. The /blog-content skill fills the JSON; this script is
pure rendering and does no research or writing.

Usage:
    python3 build_workbook.py <plan.json> <output.xlsx>

The plan JSON contract (every key optional except meta + pieces):

{
  "meta": {
    "client", "business_type", "period_label", "generated_on", "goal",
    "prepared_by", "primary_market", "semrush_used"   # bool: was live SEMRUSH data used
  },
  "brand_voice": {
    "summary", "tone_attributes": [...], "do": [...], "dont": [...],
    "reading_level", "sample_phrases": [...], "sources": [...]
  },
  "pieces": [
    {
      "id", "title", "type"("Blog"|"Landing Page"), "funnel_stage"("TOFU"|"MOFU"|"BOFU"),
      "angle", "target_date", "word_count_target", "status", "asset_id", "doc_path",
      "seo": {
        "primary_keyword": {"term","volume","kd","intent"},
        "secondary_keywords": [{"term","volume","kd"}],
        "serp_competitors": [...], "content_gap", "meta_title", "meta_description",
        "url_slug", "internal_links": [...], "schema_type", "search_intent"
      },
      "aeo": {
        "target_question", "citable_answer", "entities": [...], "structured_data",
        "eeat_signals": [...], "why_cited"
      },
      "outline": {
        "h1", "sections": [{"h2","h3":[...],"key_points":[...]}],
        "faq": [{"q","a"}], "sources": [...]
      }
    }
  ],
  "assets": [
    {"id","title","type","platform","model","settings","prompt",
     "file","thumb","web","credits"}      # file/thumb filled by download_assets.py
  ],
  "contact": {"company","tagline","contact_name","title","email","phone","website","note"}
}
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
from pathlib import Path

for _pkg in ("openpyxl", "PIL"):
    try:
        __import__(_pkg)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install",
                               "openpyxl" if _pkg == "openpyxl" else "Pillow"])

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

# ── DigitalCLIQ design tokens (Resources/design-system/Design-System.md) ──
DIGITAL_BLUE = "405FAB"   # mastheads, primary fills
SKY_BLUE = "6B9DD4"       # accents, "strong" status
RICH_BLACK = "000000"     # light-theme body text
INK = "000000"            # body text = Rich Black (legacy #0D1B2A retired)
NAVY_DEEP = "070A15"      # dark cover/contact band
WARM_GREY = "949592"      # muted labels, "weak" status
CALLOUT_TINT = "EDF2F9"   # callout fills, alternating rows
LIGHT_BLUE = "EDF2F9"     # label fills → Callout Tint
ALT_ROW = "EDF2F9"        # alternating body rows → Callout Tint
BORDER_BLUE = "D8E1F0"    # thin cell borders

HEAD_FONT = "Dosis"        # structure: headings, labels, stats
BODY_FONT = "Roboto Slab"  # reading: body cells

# Funnel-stage tint ladder, palette-only (text label TOFU/MOFU/BOFU carries meaning)
STAGE_FILL = {"TOFU": "EDF2F9", "MOFU": "D8E1F0", "BOFU": "6B9DD4"}
# Status coding, palette-only: Sky Blue = strong/done, Warm Grey = in-progress/weak,
# Digital Blue = emphasis. The status TEXT in the cell carries the meaning.
STATUS_FILL = {
    "written": SKY_BLUE, "ready": SKY_BLUE, "drafted": SKY_BLUE, "outline": WARM_GREY,
    "draft": WARM_GREY, "review": WARM_GREY, "published": DIGITAL_BLUE, "queued": WARM_GREY,
}

LOGO_PATH = Path("/Users/drewmoon/Desktop/DigitalCLIQ Brain HQ/Resources/brand-assets/"
                 "digital-cliq-logo-solid-1000px-wide.png")

# Hard guardrail: this address must NEVER appear in any DigitalCLIQ deliverable.
BANNED_EMAIL = "digitalcliq@gmail.com"
APPROVED_EMAIL = "drewmoon@digitalcliq.com"


# ── style helpers (Dosis = structure, Roboto Slab = reading) ──────────
def _f_header():  return Font(name=HEAD_FONT, size=10, bold=True, color="FFFFFF")
def _f_body(b=False): return Font(name=BODY_FONT, size=10, bold=b, color=INK)
def _f_muted():   return Font(name=BODY_FONT, size=9, italic=True, color=WARM_GREY)
def _f_title():   return Font(name=HEAD_FONT, size=20, bold=True, color="FFFFFF")
def _f_section(): return Font(name=HEAD_FONT, size=13, bold=True, color=DIGITAL_BLUE)
def _f_link():    return Font(name=BODY_FONT, size=10, underline="single", color=DIGITAL_BLUE)
def _fill(c):     return PatternFill("solid", fgColor=c)
def _border():
    s = Side(style="thin", color=BORDER_BLUE)
    return Border(left=s, right=s, top=s, bottom=s)


def _masthead(ws, title, ncols, subtitle=""):
    """Design-System §4 Excel masthead: rows 1-2 merged Digital Blue band,
    WHITE logo anchored A1 (~0.35in tall), sheet title in white Dosis.
    Optional muted subtitle lands on row 3. Fails loudly if the logo is missing."""
    if not LOGO_PATH.exists():
        raise FileNotFoundError(
            f"DigitalCLIQ logo not found at canonical path: {LOGO_PATH}. "
            "Fix the path; do not ship without branding.")
    ncols = max(int(ncols), 8)  # keep the band + title merge wide enough on narrow sheets
    for r in (1, 2):
        ws.row_dimensions[r].height = 24
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = _fill(DIGITAL_BLUE)
    img = XLImage(str(LOGO_PATH))
    lw, lh = PILImage.open(str(LOGO_PATH)).size
    img.height = 34
    img.width = int(lw * 34 / lh)
    ws.add_image(img, "A1")
    ws.merge_cells(start_row=1, start_column=3, end_row=2, end_column=ncols)
    tc = ws.cell(row=1, column=3, value=title)
    tc.font = Font(name=HEAD_FONT, size=14, bold=True, color="FFFFFF")
    tc.alignment = Alignment(horizontal="left", vertical="center")
    if subtitle:
        sc = ws.cell(row=3, column=1, value=subtitle)
        sc.font = _f_muted()
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)


def _hrow(ws, headers, row=1, fill=DIGITAL_BLUE, height=30):
    ws.row_dimensions[row].height = height
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = _f_header()
        c.fill = _fill(fill)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _border()


def _brow(ws, row_idx, values, wrap_over=28):
    fill = _fill(ALT_ROW) if row_idx % 2 == 0 else None
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=col, value=v)
        c.font = _f_body()
        c.alignment = Alignment(vertical="top",
                                wrap_text=isinstance(v, str) and len(str(v)) > wrap_over)
        c.border = _border()
        if fill:
            c.fill = fill


def _bullets(items):
    if not items:
        return ""
    if isinstance(items, str):
        return items
    return "\n".join(f"• {x}" for x in items)


def _kw(k):
    """Render a keyword dict {term, volume, kd, intent} compactly."""
    if isinstance(k, str):
        return k
    if not isinstance(k, dict):
        return str(k)
    bits = [k.get("term", "")]
    extra = []
    if k.get("volume") not in (None, ""):
        extra.append(f"vol {k['volume']}")
    if k.get("kd") not in (None, ""):
        extra.append(f"KD {k['kd']}")
    if k.get("intent"):
        extra.append(str(k["intent"]))
    if extra:
        bits.append("(" + ", ".join(extra) + ")")
    return " ".join(b for b in bits if b)


def _kw_list(items):
    if not items:
        return ""
    return "\n".join(f"• {_kw(k)}" for k in items)


def _thumb_for(path: str, max_px: int = 160) -> str | None:
    if not path or not os.path.exists(path):
        return None
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return None
    try:
        im = PILImage.open(path).convert("RGB")
        im.thumbnail((max_px, max_px))
        tdir = Path(path).parent / "_xlsx_thumbs"
        tdir.mkdir(exist_ok=True)
        out = tdir / (Path(path).stem + f"_{max_px}.png")
        im.save(out, "PNG")
        return str(out)
    except Exception:
        return None


# ── Tab 1: Cover ──────────────────────────────────────────────────────
def _cover(wb, meta):
    ws = wb.create_sheet("Cover")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100

    for r in range(1, 11):
        for c in range(1, 8):
            ws.cell(row=r, column=c).fill = _fill(NAVY_DEEP)
        ws.row_dimensions[r].height = 24

    if not LOGO_PATH.exists():
        raise FileNotFoundError(
            f"DigitalCLIQ logo not found at canonical path: {LOGO_PATH}. "
            "Fix the path; do not ship without branding.")
    img = XLImage(str(LOGO_PATH))
    lw, lh = PILImage.open(str(LOGO_PATH)).size
    img.height = 54
    img.width = int(lw * 54 / lh)  # keep aspect ratio intact (Visual-QA)
    ws.add_image(img, "B2")

    ws["B5"] = "Blog & Content Plan"
    ws["B5"].font = _f_title()
    ws["B6"] = meta.get("client", "")
    ws["B6"].font = Font(name=HEAD_FONT, size=15, bold=True, color="FFFFFF")
    ws["B7"] = meta.get("business_type", "")
    ws["B7"].font = Font(name=HEAD_FONT, size=11, color=SKY_BLUE)
    ws["B8"] = meta.get("period_label", "")
    ws["B8"].font = Font(name=BODY_FONT, size=11, color="FFFFFF")

    ws["B12"] = "The Goal"
    ws["B12"].font = _f_section()
    ws["B13"] = meta.get("goal", "")
    ws["B13"].font = _f_body()
    ws["B13"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[13].height = 44

    sem = ("Live SEMRUSH keyword + competitor data" if meta.get("semrush_used")
           else "Web + SERP research (no live SEMRUSH this run)")
    notes = [
        ("How to read this workbook", ""),
        ("Content Plan", "The batch at a glance: each piece, its type, funnel stage, primary keyword, the generated hero asset, a link to the finished doc, and status."),
        ("SEO & Keyword Brief", f"Per piece: target keyword cluster, search intent, meta title/description, URL slug, SERP gap, and internal links. Data source: {sem}."),
        ("AI Search Brief", "Per piece: the exact question it answers, the citable answer block, the entities and facts to assert, and the structured data that earns an AI-engine citation (AEO / GEO)."),
        ("Article Outlines", "Per piece: the H1/H2/H3 skeleton, FAQ questions, key points, and sources. The full written body ships as a separate branded Word doc."),
        ("Magnific Prompt Library", "Paste-ready prompts with the exact Magnific model + settings, plus Claude / Gemini / Sora fallbacks for each visual."),
        ("Generated Assets", "The real hero images / video generated via Magnific, saved into the vault and linked here."),
        ("Brand Voice", "The voice profile we matched: tone attributes, do/don't, reading level, and sample phrases pulled from the brand's own footprint."),
        ("Contact", "Your DigitalCLIQ point of contact."),
    ]
    r = 15
    for label, val in notes:
        ws[f"B{r}"] = label if not val else f"{label}  —  {val}"
        ws[f"B{r}"].font = _f_body(b=True) if not val else _f_body()
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        if val:
            ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    ws[f"B{r}"] = f"Prepared by {meta.get('prepared_by', 'DigitalCLIQ — Digital Strategy & Development')}"
    ws[f"B{r}"].font = Font(name=HEAD_FONT, size=10, bold=True, color=DIGITAL_BLUE)
    ws[f"B{r+1}"] = f"Generated {meta.get('generated_on', '')}  •  {APPROVED_EMAIL}"
    ws[f"B{r+1}"].font = _f_muted()


# ── Tab 2: Content Plan ───────────────────────────────────────────────
def _content_plan(wb, pieces):
    ws = wb.create_sheet("Content Plan")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Content Plan — this batch", 10,
              "Each row is one finished piece. The hero asset is generated; the body ships as a branded Word doc (link).")

    headers = ["#", "Title", "Type", "Funnel\nStage", "Primary Keyword",
               "Words", "Target Date", "Hero Asset", "Doc", "Status"]
    _hrow(ws, headers, row=4)

    r = 5
    for i, p in enumerate(pieces, 1):
        seo = p.get("seo", {})
        pk = seo.get("primary_keyword", "")
        vals = [
            i, p.get("title", ""), p.get("type", "Blog"),
            p.get("funnel_stage", ""), _kw(pk),
            p.get("word_count_target", ""), p.get("target_date", ""),
            "",  # hero thumb (col 8)
            "",  # doc link (col 9)
            p.get("status", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=2).font = _f_body(b=True)
        # funnel tint
        fc = STAGE_FILL.get(str(p.get("funnel_stage", "")).strip().upper())
        if fc:
            ws.cell(row=r, column=4).fill = _fill(fc)
        # status tint
        sc = STATUS_FILL.get(str(p.get("status", "")).strip().lower())
        if sc:
            scc = ws.cell(row=r, column=10)
            scc.fill = _fill(sc)
            scc.font = Font(name=HEAD_FONT, size=9, bold=True, color="FFFFFF")
            scc.alignment = Alignment(horizontal="center", vertical="center")
        # hero thumbnail
        thumb = _thumb_for(p.get("_hero_thumb", ""))
        if thumb:
            ws.row_dimensions[r].height = 120
            try:
                xi = XLImage(thumb)
                w, h = PILImage.open(thumb).size
                scale = min(150 / w, 110 / h)
                xi.width, xi.height = int(w * scale), int(h * scale)
                ws.add_image(xi, f"H{r}")
            except Exception:
                pass
        else:
            ws.row_dimensions[r].height = 64
        # doc link
        doc = p.get("doc_path", "")
        if doc:
            lc = ws.cell(row=r, column=9, value="Open doc")
            lc.hyperlink = ("file://" + os.path.abspath(doc)) if os.path.exists(doc) else doc
            lc.font = _f_link()
            lc.alignment = Alignment(horizontal="center", vertical="center")
        r += 1

    widths = [4, 40, 12, 8, 30, 7, 12, 22, 12, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "B5"


# ── Tab 3: SEO & Keyword Brief ────────────────────────────────────────
def _seo_brief(wb, pieces):
    ws = wb.create_sheet("SEO & Keyword Brief")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "SEO & Keyword Brief", 9,
              "Built for Google ranking. Keyword metrics from SEMRUSH where available; otherwise SERP-research estimates (flagged).")

    headers = ["Title", "Primary Keyword\n(vol / KD / intent)", "Secondary / Cluster",
               "Search Intent", "Meta Title (<=60c)", "Meta Description (<=155c)",
               "URL Slug", "Internal Links", "SERP Gap / Angle"]
    _hrow(ws, headers, row=4)

    r = 5
    for p in pieces:
        seo = p.get("seo", {})
        vals = [
            p.get("title", ""),
            _kw(seo.get("primary_keyword", "")),
            _kw_list(seo.get("secondary_keywords")),
            seo.get("search_intent", seo.get("intent", "")),
            seo.get("meta_title", ""),
            seo.get("meta_description", ""),
            seo.get("url_slug", ""),
            _bullets(seo.get("internal_links")),
            seo.get("content_gap", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.row_dimensions[r].height = 150
        r += 1

    widths = [30, 26, 30, 16, 32, 40, 26, 30, 38]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "B5"


# ── Tab 4: AI Search Brief (AEO / GEO) ────────────────────────────────
def _aeo_brief(wb, pieces):
    ws = wb.create_sheet("AI Search Brief")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "AI Search Brief — Answer / Generative Engine Optimization", 7,
              "How each piece earns a citation inside ChatGPT, Google AI Overviews, Perplexity, and Gemini: "
              "a clean answer block, asserted entities/facts, and structured data.")

    headers = ["Title", "Question It Answers", "Citable Answer Block (40-60 words)",
               "Entities / Facts to Assert", "Structured Data", "E-E-A-T Signals",
               "Why an AI Engine Cites It"]
    _hrow(ws, headers, row=4)

    r = 5
    for p in pieces:
        aeo = p.get("aeo", {})
        vals = [
            p.get("title", ""),
            aeo.get("target_question", ""),
            aeo.get("citable_answer", ""),
            _bullets(aeo.get("entities")),
            aeo.get("structured_data", p.get("seo", {}).get("schema_type", "")),
            _bullets(aeo.get("eeat_signals")),
            aeo.get("why_cited", ""),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.row_dimensions[r].height = 160
        r += 1

    widths = [28, 34, 44, 34, 22, 32, 38]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "B5"


# ── Tab 5: Article Outlines ───────────────────────────────────────────
def _outlines(wb, pieces):
    ws = wb.create_sheet("Article Outlines")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Article Outlines — the skeleton behind each piece", 6,
              "Full written body ships as a branded Word doc. This is the structure, the FAQ, and the sources.")

    headers = ["Title", "H1", "Section Outline (H2 / H3)", "FAQ (People-Also-Ask)",
               "Key Points", "Sources"]
    _hrow(ws, headers, row=4)

    r = 5
    for p in pieces:
        o = p.get("outline", {})
        sec_lines, kp_lines = [], []
        for s in o.get("sections", []):
            sec_lines.append(f"H2 · {s.get('h2', '')}")
            for h3 in s.get("h3", []) or []:
                sec_lines.append(f"   H3 · {h3}")
            for kp in s.get("key_points", []) or []:
                kp_lines.append(f"• {kp}")
        faq_lines = [f"Q: {f.get('q', '')}" for f in o.get("faq", [])]
        vals = [
            p.get("title", ""),
            o.get("h1", p.get("title", "")),
            "\n".join(sec_lines),
            "\n".join(faq_lines),
            "\n".join(kp_lines),
            _bullets(o.get("sources")),
        ]
        _brow(ws, r, vals)
        ws.cell(row=r, column=1).font = _f_body(b=True)
        ws.row_dimensions[r].height = 220
        r += 1

    widths = [26, 34, 44, 38, 40, 34]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "B5"


# ── Tab 6: Magnific Prompt Library ────────────────────────────────────
def _prompt_library(wb, rows):
    ws = wb.create_sheet("Magnific Prompt Library")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Magnific Prompt Library — paste-ready", 8,
              "Primary path is Magnific (model + settings given). Fallback columns work in Claude, "
              "Gemini / Nano Banana, or Sora / Veo if generating outside Magnific.")

    headers = ["ID", "Asset Title", "Type", "Placement", "Magnific Model",
               "Settings", "Primary Prompt (Magnific)", "Fallback (Claude / Gemini / Sora)"]
    _hrow(ws, headers, row=4)

    r = 5
    for row in rows:
        fb = []
        if row.get("fallback_claude"): fb.append(f"Claude/Gemini image: {row['fallback_claude']}")
        if row.get("fallback_gemini"): fb.append(f"Gemini/Nano Banana: {row['fallback_gemini']}")
        if row.get("fallback_sora"): fb.append(f"Sora/Veo video: {row['fallback_sora']}")
        vals = [
            row.get("asset_id", ""), row.get("title", ""), row.get("type", ""),
            row.get("platform", ""), row.get("magnific_model", ""),
            row.get("settings", ""), row.get("prompt", ""), "\n\n".join(fb),
        ]
        _brow(ws, r, vals)
        ws.row_dimensions[r].height = 130
        r += 1

    widths = [6, 26, 10, 14, 22, 18, 56, 50]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A5"


# ── Tab 7: Generated Assets gallery ───────────────────────────────────
def _generated(wb, rows):
    ws = wb.create_sheet("Generated Assets")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Generated Assets — live from Magnific, saved to the vault", 8)

    headers = ["Preview", "ID", "Title", "Type", "Magnific Model",
               "Credits", "Vault Path", "Web Link"]
    _hrow(ws, headers, row=3)

    r = 4
    for row in rows:
        vals = [
            "", row.get("asset_id", ""), row.get("title", ""), row.get("type", ""),
            row.get("magnific_model", ""), row.get("credits", ""),
            row.get("vault_path", ""), "",
        ]
        _brow(ws, r, vals)
        thumb = _thumb_for(row.get("thumb_path", ""), max_px=200)
        if thumb:
            ws.row_dimensions[r].height = 150
            try:
                xi = XLImage(thumb)
                w, h = PILImage.open(thumb).size
                scale = min(190 / w, 140 / h)
                xi.width, xi.height = int(w * scale), int(h * scale)
                ws.add_image(xi, f"A{r}")
            except Exception:
                pass
        else:
            ws.row_dimensions[r].height = 30
        if row.get("web_url"):
            lc = ws.cell(row=r, column=8, value="Open")
            lc.hyperlink = row["web_url"]
            lc.font = _f_link()
        r += 1

    widths = [28, 6, 26, 12, 22, 9, 48, 12]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"


# ── Tab 8: Brand Voice ────────────────────────────────────────────────
def _brand_voice(wb, bv):
    ws = wb.create_sheet("Brand Voice")
    ws.sheet_view.showGridLines = False
    _masthead(ws, "Brand Voice Profile — what we matched", 4,
              "Pulled from the brand's own site, socials, reviews, and citations. Every piece was written to this voice.")

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90

    rows = [
        ("Voice summary", bv.get("summary", "")),
        ("Tone attributes", _bullets(bv.get("tone_attributes"))),
        ("Reading level", bv.get("reading_level", "")),
        ("Do", _bullets(bv.get("do"))),
        ("Don't", _bullets(bv.get("dont"))),
        ("Sample phrases (their words)", _bullets(bv.get("sample_phrases"))),
        ("Voice sources", _bullets(bv.get("sources"))),
    ]
    r = 4
    for label, val in rows:
        if not val:
            continue
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = _f_body(b=True)
        lc.alignment = Alignment(vertical="top")
        lc.fill = _fill(LIGHT_BLUE)
        lc.border = _border()
        vc = ws.cell(row=r, column=2, value=val)
        vc.font = _f_body()
        vc.alignment = Alignment(vertical="top", wrap_text=True)
        vc.border = _border()
        ws.row_dimensions[r].height = max(30, 16 * (1 + str(val).count("\n")))
        r += 1


# ── Tab 9: Contact ────────────────────────────────────────────────────
def _contact(wb, c):
    ws = wb.create_sheet("Contact")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 56

    for r in range(1, 8):
        for col in range(1, 6):
            ws.cell(row=r, column=col).fill = _fill(NAVY_DEEP)
        ws.row_dimensions[r].height = 22
    if not LOGO_PATH.exists():
        raise FileNotFoundError(
            f"DigitalCLIQ logo not found at canonical path: {LOGO_PATH}. "
            "Fix the path; do not ship without branding.")
    img = XLImage(str(LOGO_PATH))
    lw, lh = PILImage.open(str(LOGO_PATH)).size
    img.height = 50
    img.width = int(lw * 50 / lh)  # keep aspect ratio intact (Visual-QA)
    ws.add_image(img, "B2")

    ws["B5"] = c.get("company", "DigitalCLIQ")
    ws["B5"].font = Font(name=HEAD_FONT, size=16, bold=True, color="FFFFFF")
    ws["B6"] = c.get("tagline", "Digital Strategy & Development")
    ws["B6"].font = Font(name=HEAD_FONT, size=11, color=SKY_BLUE)

    email = c.get("email", APPROVED_EMAIL)
    rows = [
        ("Your Contact", c.get("contact_name", "Drew Moon")),
        ("Title", c.get("title", "President")),
        ("Email", email),
        ("Phone", c.get("phone", "")),
        ("Website", c.get("website", "digitalcliq.com")),
    ]
    r = 9
    for label, val in rows:
        if not val:
            continue
        ws[f"B{r}"] = label
        ws[f"B{r}"].font = _f_body(b=True)
        cell = ws[f"C{r}"]
        cell.value = val
        cell.font = _f_body()
        if label == "Email":
            cell.hyperlink = f"mailto:{val}"
            cell.font = _f_link()
        if label == "Website":
            cell.hyperlink = f"https://{val.lstrip('https://').lstrip('http://')}"
            cell.font = _f_link()
        r += 1

    if c.get("note"):
        r += 1
        ws[f"B{r}"] = c["note"]
        ws[f"B{r}"].font = _f_muted()
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)


# ── auto-derive prompt library + asset gallery from compact 'assets' ──
def _auto_fallbacks(a):
    t = a.get("type", "Image")
    return {
        "fallback_claude": a.get("fallback_claude",
            "Claude / Nano Banana image: same brand system + subject; strictly on-brand, no competitor brands or logos."),
        "fallback_gemini": a.get("fallback_gemini",
            "Gemini / Nano Banana: same prompt + 'photorealistic, on-brand lighting'; use a brand-safe reference for fidelity."),
        "fallback_sora": a.get("fallback_sora",
            ("Sora / Veo: animate the still — " + a.get("title", "")) if t in ("Image", "Graphic", "Hero")
            else "Sora / Veo: same shot brief, 16:9, ~5s, subtle ambient audio."),
    }


def _expand_plan(plan: dict) -> dict:
    plan = dict(plan)
    assets = plan.get("assets")
    amap = {a["id"]: a for a in assets} if assets else {}
    # Fallback: wire hero thumbnails from a pre-built generated_assets list when no
    # top-level 'assets' registry was passed.
    if not amap and plan.get("generated_assets"):
        amap = {a.get("asset_id"): {"thumb": a.get("thumb_path", "")}
                for a in plan["generated_assets"] if a.get("asset_id")}
    if assets:
        if not plan.get("prompt_library"):
            plan["prompt_library"] = [dict(
                asset_id=a["id"], title=a.get("title", ""), type=a.get("type", ""),
                platform=a.get("platform", "Blog hero"), magnific_model=a.get("model", ""),
                settings=a.get("settings", ""), prompt=a.get("prompt", ""), **_auto_fallbacks(a))
                for a in assets]
        if not plan.get("generated_assets"):
            plan["generated_assets"] = [dict(
                asset_id=a["id"], title=a.get("title", ""), type=a.get("type", ""),
                magnific_model=a.get("model", ""), credits=a.get("credits", ""),
                vault_path=a.get("file", ""), web_url=a.get("web", ""),
                thumb_path=a.get("thumb", ""), prompt=a.get("prompt", ""))
                for a in assets]
    # wire each piece's hero thumbnail from its asset_id
    for p in plan.get("pieces", []):
        aid = p.get("asset_id")
        if aid and aid in amap:
            p["_hero_thumb"] = amap[aid].get("thumb", "")
    return plan


def _validate(path: str) -> list:
    import zipfile
    issues = []
    if "/outputs/" not in os.path.abspath(path).replace("\\", "/"):
        issues.append("file is not inside outputs/")
    wb2 = load_workbook(path)
    imgs = [n for n in zipfile.ZipFile(path).namelist() if n.startswith("xl/media")]
    if len(imgs) < 1:
        issues.append("no embedded images (logo missing)")
    blob = ""
    for ws in wb2.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str):
                    blob += c.value.lower() + " "
    if BANNED_EMAIL in blob:
        issues.append(f"BANNED {BANNED_EMAIL} present")
    for tok in ("lorem ipsum", "yyyy-mm-dd", "{client}", "{title}", "tbd_", "xxxx@"):
        if tok in blob:
            issues.append(f"unfilled template token '{tok}'")
    return issues


def build(plan: dict, output_path: str) -> str:
    plan = _expand_plan(plan)
    blob = json.dumps(plan).lower()
    if BANNED_EMAIL in blob:
        raise ValueError(
            f"REFUSED: '{BANNED_EMAIL}' appears in the plan data. DigitalCLIQ "
            f"deliverables use only {APPROVED_EMAIL}. Scrub the data and rebuild.")

    wb = Workbook()
    wb.remove(wb.active)
    _cover(wb, plan.get("meta", {}))
    _content_plan(wb, plan.get("pieces", []))
    _seo_brief(wb, plan.get("pieces", []))
    _aeo_brief(wb, plan.get("pieces", []))
    _outlines(wb, plan.get("pieces", []))
    if plan.get("prompt_library"):
        _prompt_library(wb, plan["prompt_library"])
    if plan.get("generated_assets"):
        _generated(wb, plan["generated_assets"])
    if plan.get("brand_voice"):
        _brand_voice(wb, plan["brand_voice"])
    _contact(wb, plan.get("contact", {}))

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 build_workbook.py <plan.json> <output.xlsx>")
        sys.exit(2)
    plan_path, out_path = sys.argv[1], sys.argv[2]
    with open(plan_path, encoding="utf-8") as fh:
        plan_data = json.load(fh)
    saved = build(plan_data, out_path)
    print(f"Workbook written: {saved}")
    problems = _validate(saved)
    if problems:
        print("VALIDATION ISSUES:")
        for pr in problems:
            print("  -", pr)
        sys.exit(1)
    print("Validation OK: in outputs/, branded, no banned email, no placeholders.")
