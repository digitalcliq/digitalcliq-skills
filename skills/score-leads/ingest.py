#!/usr/bin/env python3
"""
Local CRM ingestion / format adapter for the score-leads skill.
Created by DigitalCLIQ.

Auto-detects the CRM export format from its header signature and normalizes it
to the canonical lead-scoring schema, ENTIRELY LOCALLY. This is what lets the
assistant score a report without ever reading the raw file into its context or
hand-writing a transform — it just calls score_leads.py --raw <file>.

Normalized schema (the rows handed to the scorer):
    source, leads, contact, contact_pct, appts, appts_pct,
    shows, shows_pct, sales, sales_pct

Supported formats (add new adapters in ADAPTERS):
    - native        : already-normalized score_leads CSV
    - vinsolutions  : "Lead Source ROI" export (Good-Leads basis)
    - momentum      : "E-Commerce Statistics" export (per-source)
    - momentum_lsr  : "Lead Source Report" paged .xls (JasperReports; no contact column)
    - tekion        : "Lead Source Report" (New/Used split; rolled up to vendors)

CLI:
    python3 ingest.py <raw.csv|xlsx> [out.csv] [--rollup|--no-rollup]
    (prints a compact summary; writes normalized CSV if out.csv given)
"""

import sys
import os
import csv
import re
import json

NORMALIZED_FIELDS = [
    "source", "leads", "contact", "contact_pct", "appts", "appts_pct",
    "shows", "shows_pct", "sales", "sales_pct",
]

_CANON_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "reference", "canonical_vendors.json")


# ----------------------------------------------------------------------------
# Small parsing helpers
# ----------------------------------------------------------------------------

def _num(val):
    """Parse a number from messy CRM cells: strips $ , % ( ) and treats () as negative."""
    s = str(val if val is not None else "").strip()
    if not s:
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    s = s.replace(",", "").replace("%", "").replace("$", "").replace("(", "").replace(")", "").strip()
    if s in ("", "-", "—"):
        return 0.0
    try:
        v = float(s)
    except ValueError:
        return 0.0
    return -v if neg else v


def _int(val):
    return int(round(_num(val)))


def _read_table(path):
    """Read a CSV or XLSX file into (headers, list-of-dicts). First sheet for XLSX."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        # Legacy BIFF workbook (openpyxl cannot read it). First sheet, row 0 = headers.
        sheets = _read_xls_sheets(path)
        rows = sheets[0][1] if sheets else []
        if not rows:
            return [], []
        headers = [str(h).strip() for h in rows[0]]
        return headers, [{headers[i]: (r[i] if i < len(r) else None)
                          for i in range(len(headers))} for r in rows[1:]]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            out.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
        return headers, out
    # CSV (utf-8-sig strips BOM that several CRMs prepend)
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        out = [dict(d) for d in reader]
    return headers, out


def _read_xls_sheets(path):
    """Read every sheet of a legacy .xls as [(sheet_name, [[cell, ...], ...])].
    Cells come back as stripped strings ("" for blanks). Needs xlrd; loud if missing."""
    try:
        import xlrd
    except ImportError as e:
        raise ValueError(
            "Legacy .xls input needs the xlrd package: python3 -m pip install --user xlrd"
        ) from e
    wb = xlrd.open_workbook(path)
    out = []
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in sh.row(r):
                v = c.value
                if isinstance(v, float) and v.is_integer():
                    v = int(v)
                vals.append(str(v).strip())
            rows.append(vals)
        out.append((sh.name, rows))
    return out


_PCT_RE = re.compile(r"^-?[\d,]*\.?\d+\s*%$")
_INT_RE = re.compile(r"^-?[\d,]+$")


def _parse_momentum_lsr(sheets):
    """Momentum CRM "Lead Source Report" (JasperReports .xls, one sheet per printed
    page). Columns per source: Leads Created, Sched Appts (# %), Appt Shows (# %),
    Unreported Shows (# %), Sales (DMS) (# %). No contact-made column exists in this
    report, so contact is 0 and meta carries contact_available=False.

    Returns (recs, info) or (None, None) when the workbook is not this report."""
    if not sheets:
        return None, None
    first = sheets[0][1]
    flat = " | ".join(" ".join(r) for r in first[:12])
    if "Lead Source Report" not in flat or "Leads Created" not in flat:
        return None, None

    info = {"report_store": None, "report_date_range": None, "report_totals": None}
    for r in first[:12]:
        cells = [c for c in r if c]
        for i, c in enumerate(cells):
            if c.startswith("Store:") and i + 1 < len(cells):
                info["report_store"] = cells[i + 1]
            if c.startswith("Date Range:"):
                rng = c.replace("Date Range:", "").strip() or (cells[i + 1] if i + 1 < len(cells) else "")
                info["report_date_range"] = rng

    by_name = {}
    order = []
    totals_pending = False
    for _name, rows in sheets:
        for r in rows:
            cells = [c for c in r if c]
            if not cells:
                continue
            if cells[0] == "Totals":
                totals_pending = True
                continue
            if totals_pending:
                nums = [c for c in cells if _INT_RE.match(c)]
                if nums:
                    if info["report_totals"] is None:
                        info["report_totals"] = {"leads": int(nums[0].replace(",", ""))}
                    elif "sales" not in info["report_totals"] and len(nums) >= 4:
                        ints = [int(n.replace(",", "")) for n in nums]
                        info["report_totals"].update(
                            {"appts": ints[0], "shows": ints[1], "unreported": ints[2], "sales": ints[3]})
                continue
            if len(cells) != 10:
                continue
            name = cells[0]
            vals = cells[1:]
            # Layout: leads, then (#, %) pairs for sched appts, shows, unreported shows, sales
            int_idx, pct_idx = (0, 1, 3, 5, 7), (2, 4, 6, 8)
            if not all(_INT_RE.match(vals[i]) for i in int_idx) or not all(_PCT_RE.match(vals[i]) for i in pct_idx):
                continue
            ints = [int(vals[i].replace(",", "")) for i in int_idx]  # leads, appts, shows, unreported, sales
            rec = by_name.get(name)
            if rec is None:
                rec = {"source": name, "leads": 0, "contact": 0, "appts": 0, "shows": 0,
                       "unreported_shows": 0, "sales": 0}
                by_name[name] = rec
                order.append(name)
            rec["leads"] += ints[0]
            rec["appts"] += ints[1]
            rec["shows"] += ints[2]
            rec["unreported_shows"] += ints[3]
            rec["sales"] += ints[4]
    recs = [by_name[n] for n in order]
    if len(recs) < 2:
        raise ValueError("Momentum Lead Source Report recognized but fewer than 2 source rows parsed.")
    return recs, info


def _get(d, *names):
    """Fetch the first matching key from a row dict, case/space-insensitive."""
    norm = {str(k).strip().lower(): v for k, v in d.items()}
    for n in names:
        v = norm.get(n.strip().lower())
        if v is not None:
            return v
    return None


# ----------------------------------------------------------------------------
# PDF text extraction + heuristic table parsing (best-effort)
# ----------------------------------------------------------------------------

def _read_pdf_text(path):
    """Extract text from a PDF, column layout preserved. Prefers pdftotext -layout,
    falls back to pdfplumber. Loud on total failure."""
    import subprocess
    for exe in ("/opt/homebrew/bin/pdftotext", "pdftotext"):
        try:
            out = subprocess.run([exe, "-layout", path, "-"],
                                 capture_output=True, text=True, timeout=60)
            if out.returncode == 0 and out.stdout.strip():
                return out.stdout
        except (OSError, subprocess.SubprocessError):
            continue
    try:
        import pdfplumber
        parts = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001 - report, don't swallow
        raise ValueError(f"Could not extract text from PDF: {e}")


def _is_numeric_token(tok):
    s = tok.replace(",", "").replace("%", "").replace("$", "").replace("(", "").replace(")", "")
    if s in ("", "-", "—"):
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


def parse_pdf_text(text):
    """Best-effort parse of a dealer CRM lead report PDF into count records.

    Handles the common single-table layout where each data line is a source label
    followed by a trailing run of numbers. Two column shapes are recognized by the
    count of trailing numerics:
      - 9 cells: leads, contact#, contact%, appts#, appt%, shows#, show%, sales#, sale%
        (we read the COUNT columns at indices 0,1,3,5,7 and let percentages re-derive)
      - 5 cells: leads, contact#, appts#, shows#, sales#
    Lines that don't fit (headers, totals, prose) are skipped. Returns a list of
    {source, leads, contact, appts, shows, sales}. Unreliable layouts yield few
    rows, which the caller treats as "unknown format" and routes to the manual path.
    """
    records = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        tokens = line.split()
        # peel the trailing run of numeric tokens
        i = len(tokens)
        while i > 0 and _is_numeric_token(tokens[i - 1]):
            i -= 1
        nums = tokens[i:]
        label = " ".join(tokens[:i]).strip(" .:-")
        if not label or label.lower() in ("total", "totals", "grand total", "sum"):
            continue
        if len(nums) >= 9:
            vals = nums[:9]
            rec = {"source": label, "leads": _int(vals[0]), "contact": _int(vals[1]),
                   "appts": _int(vals[3]), "shows": _int(vals[5]), "sales": _int(vals[7])}
        elif len(nums) == 5:
            rec = {"source": label, "leads": _int(nums[0]), "contact": _int(nums[1]),
                   "appts": _int(nums[2]), "shows": _int(nums[3]), "sales": _int(nums[4])}
        else:
            continue
        if rec["leads"] < 1 and rec["sales"] < 1:
            continue
        records.append(rec)
    return records


# ----------------------------------------------------------------------------
# Canonical vendor roll-up
# ----------------------------------------------------------------------------

def load_canonical_map(path=_CANON_PATH):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {"rules": [], "fallback": "Other / Ungrouped"}


def canonical_vendor(name, cmap):
    """Map a raw source label to its canonical vendor via the ordered rule list."""
    t = (name or "").lower()
    for rule in cmap.get("rules", []):
        for sub in rule.get("match", []):
            if sub in t:
                return rule["name"]
    return cmap.get("fallback", "Other / Ungrouped")


def _aggregate(records, rollup, cmap):
    """Given per-row count dicts {source,leads,contact,appts,shows,sales}, optionally
    roll up by canonical vendor, then derive the normalized percentage schema."""
    if rollup:
        buckets = {}
        for r in records:
            key = canonical_vendor(r["source"], cmap)
            b = buckets.setdefault(key, {"leads": 0, "contact": 0, "appts": 0, "shows": 0, "sales": 0})
            for k in ("leads", "contact", "appts", "shows", "sales"):
                b[k] += r[k]
        records = [dict(source=k, **v) for k, v in buckets.items()]

    rows = []
    for r in records:
        leads = r["leads"]
        # Momentum re-engagement attribution can log sales/appts under a source
        # with no lead counted this period; dropping those rows understates totals.
        if leads < 1 and not (r["sales"] or r["appts"] or r["shows"] or r["contact"]):
            continue
        appts = r["appts"]
        rows.append({
            "source": r["source"],
            "leads": leads,
            "contact": r["contact"],
            "contact_pct": round(r["contact"] / leads * 100, 1) if leads > 0 else 0.0,
            "appts": appts,
            "appts_pct": round(appts / leads * 100, 1) if leads > 0 else 0.0,
            "shows": r["shows"],
            "shows_pct": round(r["shows"] / appts * 100, 1) if appts > 0 else 0.0,
            "sales": r["sales"],
            "sales_pct": round(r["sales"] / leads * 100, 1) if leads > 0 else 0.0,
        })
    rows.sort(key=lambda x: -x["leads"])
    return rows


# ----------------------------------------------------------------------------
# Format detection + adapters
# ----------------------------------------------------------------------------

def detect_format(headers):
    h = {x.strip().lower() for x in headers}
    if "primary vehicle stock type" in h and "total good leads" in h:
        return "tekion"
    if "lead_provider" in h and "contact_made_cnt" in h:
        return "momentum"
    if "lead source group" in h and ("sold from leads" in h or "good leads" in h):
        return "vinsolutions"
    # Focus / e-Commerce "Source ROI" export (prospects basis; New/Used sold split)
    if "total prospects" in h and "total sold" in h and "appts scheduled" in h:
        return "focuscrm"
    if "source" in h and "sales_pct" in h:
        return "native"
    return None


def _adapt_native(rows, headers):
    out = []
    for d in rows:
        src = (_get(d, "source") or "").strip()
        if not src:
            continue
        out.append({k: d.get(k, 0) for k in NORMALIZED_FIELDS} | {"source": src})
    # native is already in the target schema; coerce numerics
    norm = []
    for d in out:
        norm.append({
            "source": d["source"],
            "leads": _int(d.get("leads")),
            "contact": _int(d.get("contact")),
            "contact_pct": _num(d.get("contact_pct")),
            "appts": _int(d.get("appts")),
            "appts_pct": _num(d.get("appts_pct")),
            "shows": _int(d.get("shows")),
            "shows_pct": _num(d.get("shows_pct")),
            "sales": _int(d.get("sales")),
            "sales_pct": _num(d.get("sales_pct")),
        })
    return norm, False  # already normalized; no roll-up, no re-derive


def _adapt_vinsolutions(rows, headers):
    """VinSolutions Lead Source ROI. Percentages are reported on GOOD leads, so we
    use Good Leads as the basis to match the dealer's own CRM headline numbers."""
    recs = []
    for d in rows:
        src = (_get(d, "Lead Source Group") or "").strip()
        if not src:
            continue
        good = _int(_get(d, "Good Leads"))
        if good < 1:
            continue
        recs.append({
            "source": src,
            "leads": good,
            "contact": _int(_get(d, "Internet Actual Contact")),
            "appts": _int(_get(d, "Appts Set")),
            "shows": _int(_get(d, "Appts Shown")),
            "sales": _int(_get(d, "Sold from Leads")),
        })
    return recs, True  # counts -> re-derive percentages (rollup off by default)


def _adapt_momentum(rows, headers):
    """Momentum E-Commerce Statistics. Counts are explicit; leads basis, per-source."""
    recs = []
    for d in rows:
        src = (_get(d, "lead_provider") or "").strip()
        if not src:
            continue
        leads = _int(_get(d, "leads"))
        rec = {
            "source": src,
            "leads": leads,
            "contact": _int(_get(d, "contact_made_cnt")),
            "appts": _int(_get(d, "appt_set_cnt")),
            "shows": _int(_get(d, "verified_show_cnt")),
            "sales": _int(_get(d, "sale_cnt")),
        }
        # Zero-lead rows can still carry re-engagement activity (sales/appts logged
        # under a source with no lead this period); dropping them understates totals.
        if leads < 1 and not (rec["sales"] or rec["appts"] or rec["shows"] or rec["contact"]):
            continue
        recs.append(rec)
    return recs, True


def _adapt_tekion(rows, headers):
    """Tekion Lead Source Report. Split by New/Used/Ungrouped and fragmented across
    sub-products, so counts are derived from the reported (Good-Leads-basis) percentages
    and rolled up to canonical vendors. Skips the blank totals row."""
    recs = []
    for d in rows:
        grp = (_get(d, "Lead Source Group") or "").strip()
        srcn = (_get(d, "Source Name") or "").strip()
        if not grp and not srcn:
            continue  # totals row
        good = _int(_get(d, "Total Good Leads"))
        if good < 1:
            continue
        eng = _num(_get(d, "Internet/OEM Leads Engaged %"))
        appt_pct = _num(_get(d, "Appointments Scheduled %"))
        shown_pct = _num(_get(d, "Appointments Scheduled Shown %"))
        appt_cnt = round(good * appt_pct / 100)
        recs.append({
            "source": (grp + " " + srcn).strip(),
            "leads": good,
            "contact": round(good * eng / 100),
            "appts": appt_cnt,
            "shows": round(appt_cnt * shown_pct / 100),
            "sales": _int(_get(d, "Sold In Time Period")),
        })
    return recs, True


def _adapt_focuscrm(rows, headers):
    """Focus / e-Commerce 'Source ROI' export. Prospects basis, per-source. No
    contact-count column exists in this report, so contact is left 0 (unknown).
    'Appts Scheduled' -> appts, 'Appts Kept' -> shows, 'Total Sold' -> sales.
    Skips the trailing 'Average' and 'Total' summary rows."""
    recs = []
    for d in rows:
        src = (_get(d, "Source") or "").strip()
        if not src or src.lower() in ("average", "total", "totals", "grand total"):
            continue
        leads = _int(_get(d, "Total Prospects"))
        if leads < 1:
            continue
        recs.append({
            "source": src,
            "leads": leads,
            "contact": 0,  # not reported by this CRM export
            "appts": _int(_get(d, "Appts Scheduled")),
            "shows": _int(_get(d, "Appts Kept")),
            "sales": _int(_get(d, "Total Sold")),
        })
    return recs, True  # counts -> re-derive percentages


ADAPTERS = {
    "native": (_adapt_native, False),       # (adapter, default_rollup)
    "vinsolutions": (_adapt_vinsolutions, False),
    "momentum": (_adapt_momentum, False),
    "tekion": (_adapt_tekion, True),
    "focuscrm": (_adapt_focuscrm, False),
}


def ingest(path, rollup=None):
    """Detect format, normalize locally, return (rows, meta).

    rollup: None = use the format's default; True/False = force.
    meta: {format, raw_rows, sources, store_leads, store_sales, store_close}
    """
    cmap = load_canonical_map()
    ext = os.path.splitext(path)[1].lower()

    # --- PDF: heuristic single-table parse (best-effort), roll-up off by default ---
    if ext == ".pdf":
        recs = parse_pdf_text(_read_pdf_text(path))
        if len(recs) < 2:
            raise ValueError(
                "PDF parsed into fewer than 2 data rows — layout not recognized. "
                "Use the manual path (pdftotext + normalize) for this PDF."
            )
        fmt = "pdf"
        use_rollup = bool(rollup)
        rows = _aggregate(recs, use_rollup, cmap)
        raw_rows_n = len(recs)
    elif ext == ".xls" and _parse_momentum_lsr(_read_xls_sheets(path))[0] is not None:
        recs, info = _parse_momentum_lsr(_read_xls_sheets(path))
        fmt = "momentum_lsr"
        use_rollup = bool(rollup)
        rows = _aggregate([{k: r[k] for k in ("source", "leads", "contact", "appts", "shows", "sales")}
                           for r in recs], use_rollup, cmap)
        raw_rows_n = len(recs)
        extra_meta = dict(info, contact_available=False)
        tot = info.get("report_totals") or {}
        if tot:
            got = {"leads": sum(r["leads"] for r in recs), "appts": sum(r["appts"] for r in recs),
                   "shows": sum(r["shows"] for r in recs), "sales": sum(r["sales"] for r in recs)}
            mism = {k: (got[k], tot[k]) for k in got if k in tot and got[k] != tot[k]}
            if mism:
                raise ValueError(f"Parsed source rows do not add up to the report's Totals row: {mism}")
            extra_meta["totals_check"] = "PASSED " + ", ".join(f"{k}={got[k]}" for k in got)
    else:
        headers, raw = _read_table(path)
        fmt = detect_format(headers)
        if fmt is None:
            raise ValueError(
                "Unrecognized CRM format. Headers were:\n  " + ", ".join(headers) +
                "\nAdd an adapter in ingest.py or pre-normalize to: " + ",".join(NORMALIZED_FIELDS)
            )
        adapter, default_rollup = ADAPTERS[fmt]
        use_rollup = default_rollup if rollup is None else rollup
        result = adapter(raw, headers)
        if fmt == "native":
            rows, _ = result
            if use_rollup:
                recs = [{"source": r["source"], "leads": r["leads"], "contact": r["contact"],
                         "appts": r["appts"], "shows": r["shows"], "sales": r["sales"]} for r in rows]
                rows = _aggregate(recs, True, cmap)
        else:
            recs, _ = result
            rows = _aggregate(recs, use_rollup, cmap)
        raw_rows_n = len(raw)

    tl = sum(r["leads"] for r in rows)
    ts = sum(r["sales"] for r in rows)
    extra_meta = locals().get("extra_meta") or {}
    meta = {
        "format": fmt,
        "raw_rows": raw_rows_n,
        "sources": len(rows),
        "rolled_up": use_rollup,
        "store_leads": tl,
        "store_sales": ts,
        "store_close": round(ts / tl * 100, 1) if tl else 0.0,
        "contact_available": True,
    }
    meta.update(extra_meta)
    return rows, meta


def write_normalized(rows, out_path):
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=NORMALIZED_FIELDS)
        w.writeheader()
        for r in rows:
            w.writerow(r)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = {a for a in sys.argv[1:] if a.startswith("--")}
    if not args:
        print("Usage: python3 ingest.py <raw.csv|xlsx> [out.csv] [--rollup|--no-rollup]",
              file=sys.stderr)
        sys.exit(1)
    rollup = True if "--rollup" in flags else (False if "--no-rollup" in flags else None)
    rows, meta = ingest(args[0], rollup=rollup)
    if len(args) > 1:
        write_normalized(rows, args[1])
    roll = "rolled up" if meta["rolled_up"] else "per-source"
    print(f"format: {meta['format']}  |  {meta['raw_rows']} raw rows -> {meta['sources']} sources ({roll})")
    print(f"store close (raw): {meta['store_close']}% ({meta['store_sales']}/{meta['store_leads']} leads)")
    if len(args) > 1:
        print(f"normalized: {args[1]}")


if __name__ == "__main__":
    main()
