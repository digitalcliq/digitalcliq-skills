#!/usr/bin/env python3
"""Regression tests for dealership-forecast-tool scripts (verify_workbook.py, preflight.py).

Run:  python3 tests/test_forecast.py          (plain script, exits 1 on any failure)

Needs only the standard library and openpyxl, so it runs on the Mac's Python 3.9
and in Cowork. Everything it writes goes to a temp dir (or $FORECAST_TEST_OUT);
nothing lands in the repo, the runtime or the vault.

Optional real-workbook cases: set FORECAST_STERLING_DIR to a folder holding a
read-only COPY of the Sterling February 2026 forecast as sterling_orig.xlsx and
its extracted facts as sterling_facts.json. Without it those cases print SKIP.
Client data never goes into this repo.
"""
import copy
import datetime
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from contextlib import redirect_stdout

sys.dont_write_bytecode = True

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.dirname(HERE)
SCRIPTS = os.path.join(SKILL_DIR, "scripts")
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, HERE)
import verify_workbook as vw  # noqa: E402
import preflight as pf  # noqa: E402
import build_sample as bs  # noqa: E402
import openpyxl  # noqa: E402

V = os.path.join(SCRIPTS, "verify_workbook.py")
OUT = os.environ.get("FORECAST_TEST_OUT") or tempfile.mkdtemp(prefix="forecast-tests-")
os.makedirs(OUT, exist_ok=True)
STERLING = os.environ.get("FORECAST_STERLING_DIR")
EM_DASH = chr(0x2014)
results = []


def ok(name, cond, detail=""):
    results.append((name, bool(cond)))
    print("{}  {}{}".format("PASS" if cond else "FAIL", name, ("  | " + detail[-1500:]) if detail and not cond else ""))


def skip(name, why):
    print("SKIP  {}  ({})".format(name, why))


def facts_fails(f):
    return vw.check_facts(f, vw.Report()).fails


def run(*args):
    p = subprocess.run([sys.executable, V] + list(args), stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                       universal_newlines=True, env=dict(os.environ, PYTHONDONTWRITEBYTECODE="1"))
    return p.returncode, p.stdout, p.stderr


def out(name):
    return os.path.join(OUT, name)


base_facts, rows, budget = bs.make_data()

# ================================================================ facts gate
ok("valid synthetic facts pass", facts_fails(base_facts) == [], "; ".join(facts_fails(base_facts)))

f = copy.deepcopy(base_facts)
f["history_totals"]["2024"]["leads"] += 1
fl = facts_fails(f)
ok("subtotal off by 1 fails", any("history_totals[2024].leads" in m for m in fl), "; ".join(fl))

f = copy.deepcopy(base_facts)
f["scenarios"]["scenario_planner"]["+25%"]["leads"] += 1000
fl = facts_fails(f)
ok("Exec vs Scenario Planner baseline mismatch fails", any("different baselines" in m for m in fl), "; ".join(fl))

f = copy.deepcopy(base_facts)
f["sources"].append({"name": "Credit App", "leads": 100, "sales": 30, "dups": 0, "invalid": 0, "included": True})
fl = facts_fails(f)
ok("Credit App at 30% close, unconfirmed, fails", any("looks operational" in m for m in fl), "; ".join(fl))
f["sources"][-1]["confirmed_by_drew"] = "keep (test)"
ok("same source passes once confirmed", not any("looks operational" in m for m in facts_fails(f)))
f["sources"][-1]["confirmed_by_drew"] = None
f["sources"][-1]["included"] = False
ok("same source passes once excluded", not any("looks operational" in m for m in facts_fails(f)))

f = copy.deepcopy(base_facts)
f["sources"][0]["included"] = True
fl = facts_fails(f)
ok("excluded source marked included fails", any("excluded in definitions" in m for m in fl), "; ".join(fl))

f = copy.deepcopy(base_facts)
f["vendors"][0]["tier"] = "TIER 2"
fl = facts_fails(f)
ok("wrong vendor tier fails", any("rule says TIER 1" in m for m in fl), "; ".join(fl))

f = copy.deepcopy(base_facts)
f["vendors"][0], f["vendors"][2] = f["vendors"][2], f["vendors"][0]
fl = facts_fails(f)
ok("vendors out of CPS order fail", any("not sorted by cost per sale" in m for m in fl), "; ".join(fl))

f = copy.deepcopy(base_facts)
f["labels"]["net_used"] = True
ok("'(net)' without co-op fails", any("net_used" in m for m in facts_fails(f)))

f = copy.deepcopy(base_facts)
f["ads"] = {"monthly": [{"clicks": 100, "impressions": 1000, "cost": 500}, {"clicks": 10, "impressions": 4000, "cost": 100}]}
f["ads"]["ctr"] = (0.10 + 0.0025) / 2   # average of monthly ratios (wrong; pooled is 110/5000)
f["ads"]["cpc"] = 600 / 110.0
fl = facts_fails(f)
ok("averaged monthly CTR fails; pooled CPC passes", any("ads.ctr" in m for m in fl) and not any("ads.cpc" in m for m in fl),
   "; ".join(fl))

f = copy.deepcopy(base_facts)
f["elasticity"]["label"] = "Measured"
ok("'Measured' elasticity label fails", any("label" in m for m in facts_fails(f)))

f = copy.deepcopy(base_facts)
f["forecast"]["interval_method"] = "fixed 12% 80% confidence interval"
ok("fixed band called a confidence interval fails", any("confidence interval" in m for m in facts_fails(f)))
f["forecast"]["interval_method"] = "fixed +/-12% planning range, not a confidence interval"
ok("fixed band with a 'not a confidence interval' disclaimer passes", facts_fails(f) == [], "; ".join(facts_fails(f)))

f = copy.deepcopy(base_facts)
f["labels"]["forecast_subtitle"] += " | Seasonal prior is a DigitalCLIQ estimate, not NADA data"
ok("subtitle disclaiming NADA passes", facts_fails(f) == [], "; ".join(facts_fails(f)))
f["labels"]["forecast_subtitle"] += " | NADA seasonal 35%"
ok("subtitle crediting NADA fails", any("NADA" in m for m in facts_fails(f)))

f = copy.deepcopy(base_facts)
f["manifest"] = []
ok("empty reviewer manifest fails", any("manifest" in m for m in facts_fails(f)))

# keys the sidecar saves are required by the facts gate (Step 7), not first at Step 10
for path, label in ((("client", "name"), "client.name"), (("definitions", "close_rate_basis"), "close_rate_basis"),
                    (("forecast", "months"), "forecast.months")):
    f = copy.deepcopy(base_facts)
    del f[path[0]][path[1]]
    fl = facts_fails(f)
    ok("facts without {} fail at the facts gate".format(label), any(label.split(".")[-1] in m for m in fl), "; ".join(fl))
f = copy.deepcopy(base_facts)
f["forecast"]["months"][3] = "April 2026"
ok("forecast month not YYYY-MM fails", any("YYYY-MM" in m for m in facts_fails(f)))

# two-model run: honest title passes, 3-model claim fails
f2, rows2, budget2 = bs.make_data(models_ran=("holt_winters", "seasonal_prior"))
ok("2-model facts pass with a 2-model title", facts_fails(f2) == [], "; ".join(facts_fails(f2)))
ok("2-model title drops '3-Model Ensemble'", "3-Model" not in f2["labels"]["forecast_title"], f2["labels"]["forecast_title"])
ok("facts list only the models that ran", sorted(f2["models"]["monthly"]) == ["holt_winters", "seasonal_prior"])
f3 = copy.deepcopy(f2)
f3["labels"]["forecast_title"] = "2026 Lead Forecast: 3-Model Ensemble"
ok("3-model title with 2 models fails", any("3-model" in m for m in facts_fails(f3)))

# ================================================================ tier rules
t = vw.tier_rule(6500, 4, 0)
ok("$6,500 / 4 leads / 0 sales -> TIER 4 REVIEW, note", t[:3] == ("TIER 4", "REVIEW", "0 sales on $6,500"), str(t))
ok("$300 / 2 / 0 -> NO SALES YET", vw.tier_rule(300, 2, 0)[0] == "NO SALES YET")
ok("$3,380 / 0 / 0 -> UNMEASURED", vw.tier_rule(3380, 0, 0)[:3] == ("UNMEASURED", "UNMEASURED", "No CRM leads"))
ok("$0 line among costed vendors -> NO SPEND", vw.tier_rule(0, 50, 10, True)[0] == "NO SPEND")
ok("CPS is None (shown '-') for zero-sale rows", vw.tier_rule(6500, 4, 0)[3] is None)
ok("cut-offs: $499 STAR, $500 GOOD, $999 GOOD, $1,000 AVG, $1,500 REVIEW",
   [vw.tier_rule(c, 10, 1)[1] for c in (499, 500, 999, 1000, 1500)] == ["STAR", "GOOD", "GOOD", "AVG", "REVIEW"])
cut = [v["vendor"] for v in vw.cut_order(base_facts["vendors"])]
ok("AutoTrader ($6,500, 0 sales) first on the CUT list", cut and cut[0] == "AutoTrader", str(cut))

# ================================================================ model audit
f = copy.deepcopy(base_facts)
f["models"]["monthly"]["prophet"] = [x * 1.25 / 0.97 for x in f["models"]["monthly"]["prophet"]]
f["confirmations"] = {}
p = out("divergence_facts.json")
json.dump(f, open(p, "w"))
code, o, e = run("audit", p)
ok("model 25% off the ensemble triggers the divergence prompt (exit 3)", code == 3 and "[divergence]" in o, o)
f["confirmations"]["divergence"] = "Drew: accept weights (test)"
f["confirmations"]["sanity"] = "Drew: yes (test)"
json.dump(f, open(p, "w"))
code, o, e = run("audit", p)
ok("audit exits 0 once confirmed", code == 0, o)

# ================================================================ workbook gate, passing path
X = out("DigitalCLIQ_TEST_Lead_Forecast_2026-09-26")
bs.build(base_facts, rows, budget, X + ".xlsx")
json.dump(base_facts, open(X + ".facts.json", "w"), indent=1)
if os.path.exists(X + ".forecast.json"):
    os.remove(X + ".forecast.json")
code, o, e = run("check", X + ".xlsx", "--facts", X + ".facts.json", "--write-sidecar")
ok("new-spec workbook passes the numbers gate", code == 0, o + e)
side = json.load(open(X + ".forecast.json")) if os.path.exists(X + ".forecast.json") else {}
empty = [k for k, v in side.items() if v in (None, "", [], {}) and k != "excluded_sources"]
ok("sidecar JSON written, every key non-empty", side and not empty, str(empty))
ok("sidecar nested elasticity and models filled",
   side and all(side["elasticity"].get(k) not in (None, "") for k in ("slope", "p", "n", "label"))
   and side["models"].get("ran"), json.dumps(side.get("elasticity")))
wb = openpyxl.load_workbook(X + ".xlsx")
line = next((c.value for row in wb["Executive Summary"].iter_rows() for c in row
             if isinstance(c.value, str) and c.value.startswith("Lead definition:")), "")
ok("Executive Summary definition line matches the sidecar",
   side and side["report_name"] in line and side["lead_definition_column"] in line
   and side["data_window"]["start"] in line and side["data_window"]["end"] in line
   and all(s in line for s in side["excluded_sources"]), line)

# sidecar refuses shallow facts even if a caller skips the facts gate
f = copy.deepcopy(base_facts)
f["elasticity"]["leads"]["slope"] = None
f["models"]["ran"] = []
_p, _d, empty = vw.sidecar(f, X + ".xlsx")
ok("sidecar() reports empty elasticity.slope and models.ran", "elasticity.slope" in empty and "models.ran" in empty,
   str(empty))

X2 = out("two_model")
bs.build(f2, rows2, budget2, X2 + ".xlsx")
json.dump(f2, open(X2 + ".facts.json", "w"))
code, o, e = run("check", X2 + ".xlsx", "--facts", X2 + ".facts.json")
ok("2-model workbook with a 2-model title passes", code == 0, o + e)
bs.build(f2, rows2, budget2, X2 + "_bad.xlsx", forecast_title="2026 Lead Forecast: 3-Model Ensemble")
code, o, e = run("check", X2 + "_bad.xlsx", "--facts", X2 + ".facts.json")
ok("workbook claiming 3-Model with 2 models fails", code == 1 and "3-model" in o, o)


# ================================================================ workbook gate, mutations
def mutate(name, fn):
    wb = openpyxl.load_workbook(X + ".xlsx")
    fn(wb)
    path = out(name + ".xlsx")
    wb.save(path)
    return path


def check(path, with_facts=True):
    args = ["check", path] + (["--facts", X + ".facts.json"] if with_facts else [])
    return run(*args)


def m_typed_total_and_rating(wb):
    wb["24-Month History"]["C17"] = 99999
    wb["Vendor ROI"]["K5"] = "GOOD"


code, o, e = check(mutate("mutated", m_typed_total_and_rating), with_facts=False)
ok("typed TOTAL and wrong rating are caught", code == 1 and "typed number 99,999" in o and "'Vendor ROI'!K5" in o, o)


def months_to_dates(wb):
    hs = wb["24-Month History"]
    for r in range(5, hs.max_row + 1):
        v = hs.cell(r, 1).value
        if isinstance(v, str) and not v.upper().endswith("TOTAL"):
            mon, yr = v.split()
            hs.cell(r, 1).value = datetime.datetime(int(yr), bs.MON.index(mon) + 1, 1)
            hs.cell(r, 1).number_format = "mmm yyyy"


code, o, e = check(mutate("dates_ok", months_to_dates))
ok("date-typed Month labels with correct SUM totals still pass", code == 0, o + e)


def dates_and_typed_total(wb):
    months_to_dates(wb)
    wb["24-Month History"]["C17"] = 12345


p = mutate("dates_typed", dates_and_typed_total)
code, o, e = check(p, with_facts=False)
ok("date-typed months: typed TOTAL is caught without facts",
   code == 1 and "'24-Month History'!C17 is a typed number 12,345" in o, o)
code, o, e = check(p)
ok("date-typed months: typed TOTAL is caught against facts.history_totals",
   code == 1 and "facts.history_totals say" in o, o)


def title_case_total(wb):
    hs = wb["24-Month History"]
    hs["A17"] = "2024 Total"
    hs["B17"] = 1


code, o, e = check(mutate("title_case_total", title_case_total), with_facts=False)
ok("'2024 Total' in title case is still checked (typed FAIL) and warned about capitals",
   code == 1 and "'24-Month History'!B17 is a typed number 1 " in o and "not in capitals" in e, o + e)


def no_year_total(wb):
    hs = wb["24-Month History"]
    for r in range(1, hs.max_row + 1):
        v = hs.cell(r, 1).value
        if isinstance(v, str) and v.upper().endswith("TOTAL"):
            hs.cell(r, 1).value = "Year sum"


code, o, e = check(mutate("no_year_total", no_year_total))
ok("History without a '{YEAR} TOTAL' row fails", code == 1 and "no '{YEAR} TOTAL' row" in o, o)


def unlabeled_details(wb):
    hs = wb["24-Month History"]
    for r in range(5, 17):
        hs.cell(r, 1).value = None
    hs["C17"] = 12345


code, o, e = check(mutate("unlabeled_details", unlabeled_details), with_facts=False)
ok("TOTAL over unlabeled detail rows fails loudly", code == 1 and "cannot be tied to its detail rows" in o, o)


def history_literal_changed(wb):
    wb["24-Month History"]["C5"] = wb["24-Month History"]["C5"].value + 100


code, o, e = check(mutate("history_vs_facts", history_literal_changed))
ok("History TOTAL (a correct SUM) that differs from facts.history_totals fails",
   code == 1 and "2024 TOTAL leads" in o and "facts.history_totals say" in o, o)


def typed_formula_only_totals(wb):
    wb["2026 Forecast"]["E17"] = 9000        # Planning Low (-12%) TOTAL
    wb["2026 Forecast"]["H17"] = 11000       # +25% Budget TOTAL
    bd = wb["2025 Budget Detail"]
    tr = next(r for r in range(1, bd.max_row + 1) if bd.cell(r, 1).value == "TOTAL")
    bd.cell(tr, 15).value = 123456           # Annual Total column TOTAL


code, o, e = check(mutate("typed_formula_only", typed_formula_only_totals), with_facts=False)
ok("typed TOTAL in formula-only columns (Planning Low, +25%, Annual Total) all fail",
   code == 1 and "'2026 Forecast'!E17 is a typed number" in o and "'2026 Forecast'!H17 is a typed number" in o
   and "'2025 Budget Detail'!O" in o, o)


def round_total_ok(wb):
    wb["2026 Forecast"]["E17"] = "=ROUND(D17*0.88,0)"


code, o, e = check(mutate("round_total_ok", round_total_ok), with_facts=False)
ok("Planning Low TOTAL as =ROUND(D17*0.88,0) passes (within per-row rounding)", code == 0, o + e)


def round_total_bad(wb):
    wb["2026 Forecast"]["E17"] = "=ROUND(D17*0.80,0)"


code, o, e = check(mutate("round_total_bad", round_total_bad), with_facts=False)
ok("Planning Low TOTAL that does not match its rows fails", code == 1 and "'2026 Forecast'!E17 =" in o, o)


def typed_ratio_total(wb):
    wb["24-Month History"]["H17"] = 0.07     # Close Rate TOTAL typed


code, o, e = check(mutate("typed_ratio_total", typed_ratio_total), with_facts=False)
ok("typed number in a non-additive TOTAL cell (Close Rate) fails",
   code == 1 and "'24-Month History'!H17 is a typed number" in o and "non-additive" in o, o)


def honest_disclaimers(wb):
    wb["Scenario Planner"]["A23"] = ("Elasticity is estimated, not measured. There is no sales elasticity: "
                                     "sales follow leads at the 2025 close rate.")
    wb["2026 Forecast"]["A3"] = ("Prophet (25%) + Holt-Winters (40%) + Seasonal prior (DigitalCLIQ estimate) (35%) "
                                 "| 24 months training data")
    wb["2026 Forecast"]["A19"] = "Planning range, not a confidence interval."


code, o, e = check(mutate("honest_labels", honest_disclaimers))
ok("honest disclaimers and model weights in parentheses pass with --facts", code == 0, o + e)


def dishonest_labels(wb):
    wb["Scenario Planner"]["A9"] = "Measured Spend Elasticity (leads)"
    wb["Scenario Planner"]["A24"] = "Measured Spend Elasticity (sales)"
    wb["2026 Forecast"]["E4"] = "Lower (80%)"
    wb["2026 Forecast"]["G4"] = "NADA Index"


code, o, e = check(mutate("dishonest_labels", dishonest_labels))
ok("Measured / sales elasticity / 'Lower (80%)' / NADA Index labels fail with --facts",
   code == 1 and "labelled 'Measured'" in o and "sales elasticity row" in o
   and "'2026 Forecast'!E4" in o and "NADA" in o, o)

# a negation never reaches across a sentence break to hide a real claim
ok("strip_negated keeps '80% CI' after 'not tuned.'", "80% CI" in vw.strip_negated("Weights not tuned. 80% CI"))
ok("strip_negated still strips 'not a 12.5% interval'", "interval" not in vw.strip_negated("not a 12.5% interval"))


def ci_after_sentence_break(wb):
    wb["2026 Forecast"]["A19"] = "Weights not tuned. 80% CI"


code, o, e = check(mutate("ci_after_break", ci_after_sentence_break))
ok("'Weights not tuned. 80% CI' on the Forecast tab still fails", code == 1 and "'2026 Forecast'!A19" in o, o)
f = copy.deepcopy(base_facts)
f["forecast"]["interval_method"] = "fixed +/-12% band; weights not tuned. 80% CI"
ok("facts interval_method 'not tuned. 80% CI' on a fixed band fails",
   any("confidence interval" in m for m in facts_fails(f)), "; ".join(facts_fails(f)))


# ---- TOTAL rows behind a blank spacer row, under other TOTAL rows, or labelled 'TOTAL:'
def forecast_spacer(wb):
    """Move the Forecast TOTAL row and everything below it down one row; row 17 is left blank."""
    fc = wb["2026 Forecast"]
    for r in range(fc.max_row, 16, -1):
        for c in range(1, fc.max_column + 1):
            v = fc.cell(r, c).value
            if vw.is_formula(v):
                v = re.sub(r"([A-Z])(\d+)", lambda m: m.group(1) + str(int(m.group(2)) + 1)
                           if int(m.group(2)) >= 17 else m.group(0), v)
            fc.cell(r + 1, c).value = v
            fc.cell(r + 1, c)._style = copy.copy(fc.cell(r, c)._style)
            fc.cell(r, c).value = None
    wb["Scenario Planner"]["B6"] = "='2026 Forecast'!D18"


code, o, e = check(mutate("spacer_ok", forecast_spacer))
ok("Forecast TOTAL behind a blank spacer row with correct =SUM() passes", code == 0, o + e)


def spacer_sum_through_blank(wb):
    forecast_spacer(wb)
    wb["2026 Forecast"]["D18"] = "=SUM(D5:D17)"


code, o, e = check(mutate("spacer_sum_through_blank", spacer_sum_through_blank))
ok("spaced TOTAL whose =SUM() runs through the blank row passes", code == 0, o + e)


def spacer_typed_planning(wb):
    forecast_spacer(wb)
    wb["2026 Forecast"]["E18"] = 1


code, o, e = check(mutate("spacer_typed_planning", spacer_typed_planning))
ok("spacer plus a typed Planning Low TOTAL fails", code == 1 and "'2026 Forecast'!E18 is a typed number 1 " in o, o)


def spacer_short_sum(wb):
    forecast_spacer(wb)
    wb["2026 Forecast"]["E18"] = "=SUM(E5:E15)"


code, o, e = check(mutate("spacer_short_sum", spacer_short_sum))
ok("spacer plus =SUM(E5:E15) fails", code == 1 and "'2026 Forecast'!E18 is =SUM(E5:E15) but its detail rows are E5:E16"
   in o, o)


def budget_spacer_typed(wb):
    bd = wb["2025 Budget Detail"]
    tr = next(r for r in range(1, bd.max_row + 1) if bd.cell(r, 1).value == "TOTAL")
    for c in range(1, bd.max_column + 1):
        bd.cell(tr + 1, c).value = bd.cell(tr, c).value
        bd.cell(tr + 1, c)._style = copy.copy(bd.cell(tr, c)._style)
        bd.cell(tr, c).value = None
    bd.cell(tr + 1, 3).value = 999999


p = mutate("budget_spacer_typed", budget_spacer_typed)
bd_t = openpyxl.load_workbook(p)["2025 Budget Detail"]
tr1 = next(r for r in range(1, bd_t.max_row + 1) if bd_t.cell(r, 1).value == "TOTAL")
code, o, e = check(p, with_facts=False)
ok("spacer plus a typed Budget Detail monthly TOTAL fails",
   code == 1 and "'2025 Budget Detail'!C{} is a typed number 999,999".format(tr1) in o, o)


def grand_total(formula=None, typed=None, spacer=False):
    def fn(wb):
        hs = wb["24-Month History"]
        r = 32 if spacer else 31
        put = bs.put
        put(hs, r, 1, "24-MONTH TOTAL", bs.total_font, fill=bs.total_fill)
        for c in (2, 3, 4, 5, 7):
            col = chr(64 + c)
            put(hs, r, c, "={0}17+{0}30".format(col), bs.total_font, bs.NUM, bs.total_fill)
        put(hs, r, 6, "=C{0}-D{0}-E{0}".format(r), bs.total_font, bs.NUM, bs.total_fill)
        put(hs, r, 8, '=IF(F{0}>0,G{0}/F{0},"-")'.format(r), bs.total_font, bs.PCT, bs.total_fill)
        for coord, v in (formula or {}).items():
            hs[coord] = v
        for coord, v in (typed or {}).items():
            hs[coord] = v
    return fn


code, o, e = check(mutate("grand_total_ok", grand_total()))
ok("grand total under '2025 TOTAL' that adds the year TOTALs passes", code == 0, o + e)
code, o, e = check(mutate("grand_total_typed", grand_total(typed={"C31": 12345})))
ok("typed grand total under '2025 TOTAL' fails",
   code == 1 and "'24-Month History'!C31 is a typed number 12,345" in o and "rows 17, 30" in o, o)
code, o, e = check(mutate("grand_total_wrong", grand_total(formula={"C31": "=C30"})))
ok("grand total that does not add the year TOTALs fails",
   code == 1 and "'24-Month History'!C31 =" in o and "TOTAL rows above it (rows 17, 30)" in o, o)
code, o, e = check(mutate("grand_total_spacer_double", grand_total(formula={"C32": "=SUM(C5:C30)"}, spacer=True)))
ok("grand total behind a spacer that double-counts a year TOTAL fails",
   code == 1 and "'24-Month History'!C32 =" in o and "rows 17, 30" in o, o)


def colon_label(typed=False):
    def fn(wb):
        wb["2026 Forecast"]["A17"] = "TOTAL:"
        if typed:
            wb["2026 Forecast"]["E17"] = 9000
    return fn


code, o, e = check(mutate("colon_ok", colon_label()))
ok("'TOTAL:' label with correct =SUM() passes", code == 0, o + e)
code, o, e = check(mutate("colon_typed", colon_label(typed=True)))
ok("'TOTAL:' label is a TOTAL row: a typed Planning Low TOTAL fails",
   code == 1 and "'2026 Forecast'!E17 is a typed number 9,000" in o, o)


# ---- Budget Detail TOTAL = History spend TOTAL for that year (and facts.history_totals)
def budget_bump(wb):
    bd = wb["2025 Budget Detail"]
    bd["C5"] = bd["C5"].value + 500      # double-counts part of one vendor's January


def budget_drop_line(wb):
    bd = wb["2025 Budget Detail"]
    for c in range(3, 15):
        bd.cell(6, c).value = 0          # a vendor line dropped from the budget


code, o, e = check(mutate("budget_bump", budget_bump), with_facts=False)
ok("Budget Detail TOTAL that differs from the History 2025 TOTAL spend fails",
   code == 1 and "(Budget Detail TOTAL) =" in o and "2025 TOTAL spend" in o, o)
code, o, e = check(out("budget_bump.xlsx"))
ok("... and fails against facts.history_totals[2025].spend", code == 1 and "facts.history_totals[2025].spend" in o, o)
code, o, e = check(mutate("budget_drop_line", budget_drop_line))
ok("Budget Detail with a vendor line dropped fails", code == 1 and "2025 TOTAL spend" in o, o)


def budget_no_total(wb):
    bd = wb["2025 Budget Detail"]
    tr = next(r for r in range(1, bd.max_row + 1) if bd.cell(r, 1).value == "TOTAL")
    for c in range(1, bd.max_column + 1):
        bd.cell(tr, c).value = None


code, o, e = check(mutate("budget_no_total", budget_no_total))
ok("Budget Detail without a TOTAL row warns and still ties its rows to the History spend",
   code == 0 and "has no TOTAL row" in e, o + e)

# ================================================================ design checks on the build
PAL = set(bs.PALETTE)
PALETTE = {"FF" + c for c in PAL} | {"00" + c for c in PAL} | PAL | {"00000000"}
wb = openpyxl.load_workbook(X + ".xlsx")
bad_font, bad_fill, dashes = [], [], []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.font.name not in ("Dosis", "Roboto Slab"):
                bad_font.append("{}!{} {}".format(ws.title, c.coordinate, c.font.name))
            fg = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
            if fg and fg not in PALETTE:
                bad_fill.append("{}!{} {}".format(ws.title, c.coordinate, fg))
            if isinstance(c.value, str) and EM_DASH in c.value:
                dashes.append("{}!{}".format(ws.title, c.coordinate))
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.dxf and rule.dxf.fill is not None and rule.dxf.fill.bgColor is not None:
                col = rule.dxf.fill.bgColor.rgb
                if col and col not in PALETTE:
                    bad_fill.append("{} cf {}".format(ws.title, col))
            if rule.dxf and rule.dxf.font is not None and rule.dxf.font.color is not None:
                col = rule.dxf.font.color.rgb
                if col and col not in PALETTE:
                    bad_fill.append("{} cf font {}".format(ws.title, col))
ok("every non-empty cell is Dosis or Roboto Slab", not bad_font, "; ".join(bad_font[:5]))
ok("no fill outside the palette (incl. conditional formats)", not bad_fill, "; ".join(bad_fill[:5]))
ok("no em dash in any cell", not dashes, "; ".join(dashes[:5]))
ok("tab colors from the palette", all((ws.sheet_properties.tabColor.rgb or "")[-6:] in PAL for ws in wb.worksheets))
if bs.LOGO:
    ok("logo image on every tab", all(len(ws._images) == 1 for ws in wb.worksheets),
       str([(ws.title, len(ws._images)) for ws in wb.worksheets]))
    logo_bytes = open(bs.LOGO, "rb").read()
    with zipfile.ZipFile(X + ".xlsx") as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/")]
        ok("xl/media holds an image the size of the white logo", media and any(len(z.read(n)) > 1000 for n in media),
           str(media[:3]))
else:
    skip("logo checks", "white logo not found; set DIGITALCLIQ_VAULT_ROOT")

# ================================================================ Sterling copy (optional)
if STERLING and os.path.isfile(os.path.join(STERLING, "sterling_orig.xlsx")):
    s_x = os.path.join(STERLING, "sterling_orig.xlsx")
    s_f = os.path.join(STERLING, "sterling_facts.json")
    local = out("sterling_copy.xlsx")
    shutil.copyfile(s_x, local)  # never read the original in place
    code, o, e = run("check", local)
    ok("Sterling Feb 2026 copy fails the gate (exit 1)", code == 1, o)
    ok("Sterling: every formula recomputes to its saved value", "disagree with their saved values" not in o
       and "recomputed 440 formula cells" in o, o[:400])
    for needle in ("required tab missing: Reallocation Strategy", "'24-Month History'!H14 is a typed number",
                   "'24-Month History'!I14 is a typed number", "has no Tier column",
                   "not sorted by cost per sale", "different baselines", "'(net)' label", "em dash"):
        ok("Sterling flags: " + needle, needle in o, o)
    ok("Sterling: Executive Summary D12 and C16:F16 are not flagged as Actual-column reads",
       not re.search(r"'Executive Summary'!(?:D12|[C-F]16) \S+ sits in the", o), o)
    if os.path.isfile(s_f):
        sf = json.load(open(s_f))
        el = vw.elasticity_from_history(sf["history"])
        ok("Sterling elasticity slope about 0.060, p about 0.64, CI crosses zero",
           abs(el["slope"] - 0.060) < 0.001 and abs(el["p"] - 0.643) < 0.01 and el["ci95"][0] < 0 < el["ci95"][1],
           json.dumps({k: el[k] for k in ("slope", "p", "ci95")}))
        ok("Sterling elasticity label is 'Estimated, not statistically significant'",
           el["label"] == "Estimated, not statistically significant")
        code, o, e = run("audit", s_f)
        ok("Sterling audit asks Drew (exit 3, reverses direction)", code == 3 and "reverses direction" in o, o)
        by = {v["vendor"]: v for v in sf.get("vendors") or []}
        want = {"Gubagoo Virtual Retailing": "STAR", "Edmunds": "AVG", "CARFAX": "AVG", "Costco Auto Program": "GOOD"}
        got = {n: vw.tier_rule(by[n]["spend"], by[n]["leads"], by[n]["sales"])[1] for n in want if n in by}
        ok("Sterling re-tier (Gubagoo VR STAR, Edmunds AVG, CARFAX AVG, Costco GOOD)", got == want, str(got))
    # stale cached values: edit one literal in the XML without recalculating
    stale = out("sterling_stale.xlsx")
    with zipfile.ZipFile(local) as zin, zipfile.ZipFile(stale, "w", zipfile.ZIP_DEFLATED) as zout:
        edited = False
        for n in zin.namelist():
            data = zin.read(n)
            if n == "xl/worksheets/sheet2.xml":
                s = data.decode()
                s2 = re.sub(r'(<c r="C2"[^>]*>)<v>(\d+)(?:\.0)?</v>', lambda m: m.group(1) + "<v>{}</v>".format(
                    int(m.group(2)) + 100), s, count=1)
                edited = s2 != s
                data = s2.encode()
            zout.writestr(n, data)
    code, o, e = run("check", stale)
    ok("stale saved values are caught (recalc not run)",
       edited and code == 1 and "disagree with their saved values" in o, o[-600:])
else:
    skip("Sterling real-workbook cases", "set FORECAST_STERLING_DIR to a folder with sterling_orig.xlsx")

# ================================================================ preflight (fully faked, host-independent)


class FakeSys(object):
    def __init__(self, ver):
        self.version_info = ver
        self.version = "{}.{}.{} (fake)".format(*ver)
        self.argv = []


fake_skills = out("fake_skills")
os.makedirs(os.path.join(fake_skills, "xlsx", "scripts"), exist_ok=True)
fake_recalc = os.path.join(fake_skills, "xlsx", "scripts", "recalc.py")
open(fake_recalc, "w").write("# dummy\n")
os.makedirs(os.path.join(fake_skills, "dealership-forecast-tool"), exist_ok=True)
saved_env = os.environ.pop("XLSX_RECALC", None)
ok("recalc found as the sibling xlsx skill",
   pf.find_recalc(os.path.join(fake_skills, "dealership-forecast-tool")) == fake_recalc)
vault = os.environ.get("DIGITALCLIQ_VAULT_ROOT", bs.VAULT)
orig = (pf.sys, pf.shutil.which, pf.importlib.import_module, pf.find_recalc)


def preflight(ver, soffice=True, recalc=True, missing=(), vault_path=vault):
    pf.sys = FakeSys(ver)
    pf.shutil.which = (lambda n: "/usr/bin/soffice") if soffice else (lambda n: None)
    pf.find_recalc = (lambda d: fake_recalc) if recalc else (lambda d: None)

    def fake_import(m):
        if m in missing:
            raise ImportError(m)
    pf.importlib.import_module = fake_import
    buf = io.StringIO()
    with redirect_stdout(buf):
        rc = pf.main(["--vault", vault_path])
    return rc, buf.getvalue()


try:
    rc, o = preflight((3, 9, 6), soffice=False, recalc=False)
    ok("Python 3.9, no soffice -> exit 2 with the Cowork message and no pip line",
       rc == 2 and "Run it in Cowork" in o and "--break-system-packages" not in o, o)
    rc, o = preflight((3, 12, 3), recalc=False)
    ok("capable machine without recalc.py -> exit 5 naming XLSX_RECALC, not 'Run it in Cowork'",
       rc == 5 and "XLSX_RECALC" in o and "Run it in Cowork" not in o, o)
    rc, o = preflight((3, 12, 3), missing=("prophet",))
    ok("capable machine missing prophet -> exit 3 with the Cowork-only pip line",
       rc == 3 and "--break-system-packages" in o and "Cowork only" in o, o)
    if os.path.isdir(vault):
        rc, o = preflight((3, 12, 3))
        ok("capable machine, vault reachable -> exit 0 with recalc command", rc == 0 and "recalc.py" in o, o)
    else:
        skip("preflight exit 0", "vault not found; set DIGITALCLIQ_VAULT_ROOT")
    rc, o = preflight((3, 12, 3), vault_path=out("no_such_vault"))
    ok("missing vault/logo -> exit 4", rc == 4 and "logo" in o, o)
finally:
    pf.sys, pf.shutil.which, pf.importlib.import_module, pf.find_recalc = orig
os.environ["XLSX_RECALC"] = fake_recalc
ok("XLSX_RECALC override is honored", pf.find_recalc(out("elsewhere")) == fake_recalc)
if saved_env is None:
    os.environ.pop("XLSX_RECALC", None)
else:
    os.environ["XLSX_RECALC"] = saved_env

print()
n_fail = sum(1 for r in results if not r[1])
print("{} tests, {} failed  (output in {})".format(len(results), n_fail, OUT))
sys.exit(1 if n_fail else 0)
