"""Synthetic test fixture for dealership-forecast-tool (not a deliverable generator).

Builds an 8-tab workbook to references/excel_template.md plus a consistent facts
dict, with made-up numbers for a made-up store ("Test Motors", code TEST), so
tests/test_forecast.py can exercise scripts/verify_workbook.py on the passing
path and on targeted mutations. Writes only to the path it is given.

    python3 tests/build_sample.py /tmp/out/DigitalCLIQ_TEST_Lead_Forecast.xlsx
"""
import datetime
import json
import os
import random
import sys

sys.dont_write_bytecode = True

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.dirname(HERE), "scripts"))
import verify_workbook as vw  # noqa: E402

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.drawing.image import Image as XLImage

VAULT = os.environ.get("DIGITALCLIQ_VAULT_ROOT", "/Users/drewmoon/Desktop/DigitalCLIQ Brain HQ")
WHITE_LOGO = os.path.join(VAULT, "Resources", "brand-assets", "digital-cliq-logo-solid-1000px-wide.png")
LOGO = WHITE_LOGO if os.path.isfile(WHITE_LOGO) else None  # None: numbers-only fixture, no image

DIGITAL_BLUE, SKY_BLUE, WARM_GREY, TILE_BLUE = "405FAB", "6B9DD4", "949592", "2E4780"
CALLOUT_TINT, CARD_WHITE, BORDER_BLUE, WHITE, RICH_BLACK = "EDF2F9", "FBFBFD", "D8E1F0", "FFFFFF", "000000"
PALETTE = {DIGITAL_BLUE, SKY_BLUE, WARM_GREY, TILE_BLUE, CALLOUT_TINT, CARD_WHITE, BORDER_BLUE, WHITE, RICH_BLACK}
H, B = "Dosis", "Roboto Slab"
title_font = Font(name=H, bold=True, size=14, color=WHITE)
date_font = Font(name=H, size=10, color=WHITE)
hdr_font = Font(name=H, bold=True, size=10, color=WHITE)
band_fill = PatternFill("solid", fgColor=DIGITAL_BLUE)
total_fill = PatternFill("solid", fgColor=TILE_BLUE)
total_font = Font(name=H, bold=True, size=10, color=WHITE)
label_font = Font(name=H, bold=True, size=10, color=RICH_BLACK)
body_font = Font(name=B, size=10, color=RICH_BLACK)
note_font = Font(name=B, italic=True, size=9, color=WARM_GREY)
alt_fill = PatternFill("solid", fgColor=CALLOUT_TINT)
input_font = Font(name=B, bold=True, size=10, color=DIGITAL_BLUE)
thin = Side(style="thin", color=BORDER_BLUE)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
input_border = Border(*[Side(style="medium", color=DIGITAL_BLUE)] * 4)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
CUR, CUR2, PCT, NUM, SPCT = "$#,##0", "$#,##0.00", "0.0%", "#,##0", "+0.0%;-0.0%"
MON = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()
PRIOR = [0.88, 0.94, 1.08, 1.04, 1.02, 1.00, 0.96, 1.03, 1.02, 1.05, 0.98, 1.00]


def put(ws, r, c, v, font=body_font, fmt=None, fill=None, bd=True):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if bd:
        cell.border = border
    return cell


def masthead(ws, title, ncols, color):
    ws.sheet_properties.tabColor = color
    for r in (1, 2):
        ws.row_dimensions[r].height = 22
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = band_fill
    if LOGO:  # the real skill stops when the logo is missing; the fixture just skips it
        img = XLImage(LOGO)
        img.height = 34
        img.width = int(34 * 500 / 154)
        ws.add_image(img, "A1")
    put(ws, 1, 3, title, title_font, bd=False)
    put(ws, 1, ncols, datetime.date(2026, 9, 26).isoformat(), date_font, bd=False)
    ws.cell(row=1, column=ncols).alignment = Alignment(horizontal="right")


def header(ws, row, cols):
    for i, t in enumerate(cols, 1):
        c = put(ws, row, i, t, hdr_font, fill=band_fill)
        c.alignment = ctr
    ws.freeze_panes = "A{}".format(row + 1)
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = 16
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 30


def tier_rules(ws, rng, ref):
    styles = [('"TIER 1"', PatternFill("solid", bgColor=SKY_BLUE), Font(color=WHITE, bold=True)),
              ('"TIER 2"', PatternFill("solid", bgColor=CALLOUT_TINT), Font(color=DIGITAL_BLUE, bold=True)),
              ('"TIER 4"', PatternFill("solid", bgColor=WARM_GREY), Font(color=WHITE, bold=True))]
    for val, fill, font in styles:
        rule = Rule(type="expression", dxf=DifferentialStyle(fill=fill, font=font))
        rule.formula = ["{}={}".format(ref, val)]
        ws.conditional_formatting.add(rng, rule)
    muted = Rule(type="expression", dxf=DifferentialStyle(font=Font(color=WARM_GREY, italic=True)))
    muted.formula = ['OR({0}="UNMEASURED",{0}="NO SALES YET",{0}="NO SPEND")'.format(ref)]
    ws.conditional_formatting.add(rng, muted)


# ------------------------------------------------------------------ data

def make_data(seed=7, models_ran=("prophet", "holt_winters", "seasonal_prior")):
    rnd = random.Random(seed)
    vendors_raw = [  # vendor, annual spend, leads, sales
        ("DealerInspire Website", 24000, 1400, 150),
        ("Gubagoo Virtual Retailing", 13200, 1500, 45),
        ("CarGurus", 40000, 870, 50),
        ("Cars.com", 30000, 430, 28),
        ("AutoTrader", 6500, 4, 0),
        ("Small Listing Site", 300, 2, 0),
        ("LotLinx", 3380, 0, 0),
        ("Service Department", 0, 50, 10),
    ]
    rows = []
    for name, spend, leads, sales in vendors_raw:
        tier, rating, note, cps = vw.tier_rule(spend, leads, sales, True)
        rows.append({"vendor": name, "spend": spend, "leads": leads, "sales": sales, "good": leads,
                     "tier": tier, "rating": rating, "note": note, "_cps": cps,
                     "sources": [name]})
    rows.sort(key=lambda v: (v["_cps"] is None, v["_cps"] or 0))
    # vendor monthly spend for 2025 (budget detail); history 2025 spend = monthly sums
    budget = []
    for v in rows:
        m = [round(v["spend"] / 12.0) for _ in range(12)]
        m[-1] += v["spend"] - sum(m)
        budget.append((v["vendor"], m))
    hist = []
    for y in (2024, 2025):
        for i in range(12):
            if y == 2025:
                spend = sum(b[1][i] for b in budget) + 50000 + rnd.randint(-8000, 8000)
            else:
                spend = 50000 + 110000 / 12 + rnd.randint(-9000, 9000)
            leads = int(800 * PRIOR[i] * (spend / 60000.0) ** 0.08 + rnd.randint(-25, 25))
            dups = int(leads * 0.08)
            invalid = int(leads * 0.02)
            sales = int((leads - dups - invalid) * 0.07)
            hist.append({"month": "%d-%02d" % (y, i + 1), "spend": int(spend), "leads": leads, "sales": sales,
                         "dups": dups, "invalid": invalid})
    # the 2025 budget detail must sum to 2025 history spend: put the remainder on a PPC line
    budget.append(("PPC", [0] * 12))
    ppc = budget[-1][1]
    for i in range(12):
        ppc[i] = hist[12 + i]["spend"] - sum(b[1][i] for b in budget[:-1])
    totals = {}
    for h in hist:
        t = totals.setdefault(h["month"][:4], {"spend": 0, "leads": 0, "sales": 0, "dups": 0, "invalid": 0})
        for k in t:
            t[k] += h[k]
    # models
    avg = sum(h["leads"] for h in hist) / 24.0
    growth = totals["2025"]["leads"] / float(totals["2024"]["leads"])
    monthly = {
        "seasonal_prior": [round(avg * growth * PRIOR[i], 1) for i in range(12)],
        "holt_winters": [round(hist[12 + i]["leads"] * 1.01, 1) for i in range(12)],
        "prophet": [round(hist[12 + i]["leads"] * 0.97, 1) for i in range(12)],
    }
    w_all = {"prophet": 0.25, "holt_winters": 0.40, "seasonal_prior": 0.35}
    ran = list(models_ran)
    tot = sum(w_all[m] for m in ran)
    weights = {m: round(w_all[m] / tot, 4) for m in ran}
    base = [int(round(sum(weights[m] * monthly[m][i] for m in ran))) for i in range(12)]
    el = vw.elasticity_from_history(hist)
    slope = round(el["slope"], 3)
    ratio = totals["2025"]["sales"] / float(totals["2025"]["leads"])
    L = sum(base)
    bsales = int(vw._round_half_up(L * ratio, 0))

    def scen(adj):
        leads = int(vw._round_half_up(L * (1 + adj * slope), 0))
        return {"leads": leads, "sales": int(vw._round_half_up(leads * ratio, 0))}

    scen_block = {"baseline": {"leads": L, "sales": bsales}, "+25%": scen(0.25), "-15%": scen(-0.15)}
    title = ("2026 Lead Forecast: 3-Model Ensemble" if len(ran) == 3 else
             "2026 Lead Forecast: {}-Model Blend ({})".format(len(ran), " + ".join(
                 {"prophet": "Prophet", "holt_winters": "Holt-Winters",
                  "seasonal_prior": "Seasonal prior"}[m] for m in ran)))
    names = {"prophet": "Prophet", "holt_winters": "Holt-Winters", "seasonal_prior": "Seasonal prior (DigitalCLIQ estimate)"}
    subtitle = " + ".join("{} {:.0%}".format(names[m], weights[m]) for m in ran) + \
        " | 24 months training data (Jan 2024 to Dec 2025)"
    facts = {
        "schema": "forecast-facts/1",
        "client": {"code": "TEST", "name": "Test Motors"},
        "definitions": {"crm": "VinSolutions", "report_name": "Lead Source ROI", "lead_definition_column": "total_leads",
                        "close_rate_basis": "Sales / Good Leads (total_leads - total_dups - total_invalid_leads)",
                        "excluded_sources": ["Dealer Website Credit Application"],
                        "data_window": {"start": "2024-01", "end": "2025-12"}},
        "history": hist,
        "history_totals": totals,
        "sources": [{"name": "Dealer Website Credit Application", "vendor": "Credit App", "leads": 240, "sales": 66,
                     "dups": 0, "invalid": 0, "included": False, "confirmed_by_drew": "exclude"}] +
                   [{"name": v["vendor"], "leads": v["leads"], "sales": v["sales"], "included": True,
                     "confirmed_by_drew": "keep" if v["vendor"] == "Service Department" else None} for v in rows],
        "vendors": [{k: v[k] for k in ("vendor", "spend", "leads", "sales", "tier", "rating", "note", "sources")}
                    for v in rows],
        "models": {"ran": ran, "weights": weights, "monthly": {m: monthly[m] for m in ran},
                   "versions": {"python": "3.12 (synthetic)"}, "training_window": {"start": "2024-01", "end": "2025-12"}},
        "forecast": {"months": ["2026-%02d" % (i + 1) for i in range(12)], "baseline": base,
                     "low": [int(vw._round_half_up(b * 0.88, 0)) for b in base],
                     "high": [int(vw._round_half_up(b * 1.12, 0)) for b in base],
                     "interval_method": "fixed +/-12% planning range"},
        "elasticity": {"leads": {"slope": slope, "r": round(el["r"], 3), "p": round(el["p"], 3),
                                 "se": round(el["se"], 3), "n": el["n"],
                                 "ci95": [round(el["ci95"][0], 3), round(el["ci95"][1], 3)]},
                       "label": el["label"]},
        "baseline": {"spend": totals["2025"]["spend"], "leads": L, "sales": bsales, "close_rate_year": "2025"},
        "scenarios": {"executive_summary": scen_block, "scenario_planner": json.loads(json.dumps(scen_block))},
        "coop": {"present": False, "amount": 0},
        "labels": {"forecast_title": title, "forecast_subtitle": subtitle, "net_used": False},
        "confirmations": {},
        "manifest": [{"value": "{:,}".format(L), "label": "Executive Summary, 2026 Baseline leads",
                      "source": "facts.forecast.baseline (ensemble)"}],
    }
    prompts = vw.audit_prompts(facts, vw.Report())
    for k, _p in prompts:
        facts["confirmations"][k] = "Drew: yes (synthetic test)"
    return facts, rows, budget


# ---------------------------------------------------------------- workbook

def build(facts, rows, budget, path, forecast_title=None):
    wb = Workbook()
    es = wb.active
    es.title = "Executive Summary"
    hs = wb.create_sheet("24-Month History")
    vr = wb.create_sheet("Vendor ROI")
    fc = wb.create_sheet("2026 Forecast")
    sp = wb.create_sheet("Scenario Planner")
    bd = wb.create_sheet("2025 Budget Detail")
    ap = wb.create_sheet("All Providers")
    rs = wb.create_sheet("Reallocation Strategy")
    HN, FN, SN = "'24-Month History'", "'2026 Forecast'", "'Scenario Planner'"
    defs = facts["definitions"]
    el = facts["elasticity"]
    lead_el = el["leads"]

    # ---- History
    masthead(hs, "Monthly History: Jan 2024 to Dec 2025", 11, SKY_BLUE)
    put(hs, 3, 1, "Close rate = Sales / Good Leads (Leads - Duplicates - Invalid).", note_font, bd=False)
    header(hs, 4, ["Month", "Total Spend", "Total Leads", "Duplicates", "Invalid", "Good Leads", "Total Sales",
                   "Close Rate", "Cost/Lead", "Cost/Sale", "Dup Rate"])
    r = 5
    tot_rows = {}
    for yi, y in enumerate(("2024", "2025")):
        first = r
        for h in facts["history"][yi * 12:(yi + 1) * 12]:
            fill = alt_fill if (r % 2 == 0) else None
            mm = int(h["month"][5:7])
            put(hs, r, 1, "{} {}".format(MON[mm - 1], y), label_font, fill=fill)
            put(hs, r, 2, h["spend"], fmt=CUR, fill=fill)
            put(hs, r, 3, h["leads"], fmt=NUM, fill=fill)
            put(hs, r, 4, h["dups"], fmt=NUM, fill=fill)
            put(hs, r, 5, h["invalid"], fmt=NUM, fill=fill)
            put(hs, r, 6, "=C{0}-D{0}-E{0}".format(r), fmt=NUM, fill=fill)
            put(hs, r, 7, h["sales"], fmt=NUM, fill=fill)
            put(hs, r, 8, '=IF(F{0}>0,G{0}/F{0},"-")'.format(r), fmt=PCT, fill=fill)
            put(hs, r, 9, '=IF(C{0}>0,B{0}/C{0},"-")'.format(r), fmt=CUR2, fill=fill)
            put(hs, r, 10, '=IF(G{0}>0,B{0}/G{0},"-")'.format(r), fmt=CUR, fill=fill)
            put(hs, r, 11, '=IF(C{0}>0,D{0}/C{0},"-")'.format(r), fmt=PCT, fill=fill)
            r += 1
        last = r - 1
        put(hs, r, 1, "{} TOTAL".format(y), total_font, fill=total_fill)
        for c, f in ((2, CUR), (3, NUM), (4, NUM), (5, NUM), (7, NUM)):
            col = chr(64 + c)
            put(hs, r, c, "=SUM({0}{1}:{0}{2})".format(col, first, last), total_font, f, total_fill)
        put(hs, r, 6, "=C{0}-D{0}-E{0}".format(r), total_font, NUM, total_fill)
        put(hs, r, 8, '=IF(F{0}>0,G{0}/F{0},"-")'.format(r), total_font, PCT, total_fill)
        put(hs, r, 9, '=IF(C{0}>0,B{0}/C{0},"-")'.format(r), total_font, CUR2, total_fill)
        put(hs, r, 10, '=IF(G{0}>0,B{0}/G{0},"-")'.format(r), total_font, CUR, total_fill)
        put(hs, r, 11, '=IF(C{0}>0,D{0}/C{0},"-")'.format(r), total_font, PCT, total_fill)
        tot_rows[y] = r
        r += 1
    T1, T2 = tot_rows["2024"], tot_rows["2025"]

    # ---- Vendor ROI
    masthead(vr, "Vendor ROI Analysis: 2025 Spend vs Lead Performance", 12, SKY_BLUE)
    put(vr, 3, 1, "Excludes: " + ", ".join(defs["excluded_sources"]) + " (confirmed with Drew). "
        "Tier cut-offs are a DigitalCLIQ planning assumption.", note_font, bd=False)
    header(vr, 4, ["Tier", "Vendor", "Annual Spend", "Leads", "Good Leads", "Sales", "Close Rate", "Cost/Lead",
                   "Cost/Sale", "Lead Share", "Rating", "Note"])
    f, l = 5, 4 + len(rows)
    for i, v in enumerate(rows):
        r = f + i
        fill = alt_fill if (r % 2 == 0) else None
        put(vr, r, 1, ('=IF(AND(C{r}=0,COUNTIF($C${f}:$C${l},">0")>0),"NO SPEND",IF(C{r}=0,"",'
                       'IF(D{r}=0,"UNMEASURED",IF(F{r}=0,IF(C{r}>=1500,"TIER 4","NO SALES YET"),'
                       'IF(NOT(ISNUMBER(I{r})),"UNMEASURED",IF(I{r}<500,"TIER 1",IF(I{r}<1000,"TIER 2",'
                       'IF(I{r}<1500,"TIER 3","TIER 4"))))))))').format(r=r, f=f, l=l), label_font, fill=fill)
        put(vr, r, 2, v["vendor"], label_font, fill=fill)
        put(vr, r, 3, v["spend"], fmt=CUR, fill=fill)
        put(vr, r, 4, v["leads"], fmt=NUM, fill=fill)
        put(vr, r, 5, v["good"], fmt=NUM, fill=fill)
        put(vr, r, 6, v["sales"], fmt=NUM, fill=fill)
        put(vr, r, 7, '=IF(E{0}>0,F{0}/E{0},"-")'.format(r), fmt=PCT, fill=fill)
        put(vr, r, 8, '=IF(D{0}>0,C{0}/D{0},"-")'.format(r), fmt=CUR2, fill=fill)
        put(vr, r, 9, '=IF(AND(C{0}>0,F{0}>0),C{0}/F{0},"-")'.format(r), fmt=CUR, fill=fill)
        put(vr, r, 10, "=D{0}/SUM($D${1}:$D${2})".format(r, f, l), fmt=PCT, fill=fill)
        put(vr, r, 11, '=IF(A{0}="TIER 1","STAR",IF(A{0}="TIER 2","GOOD",IF(A{0}="TIER 3","AVG",'
                       'IF(A{0}="TIER 4","REVIEW",A{0}))))'.format(r), label_font, fill=fill)
        put(vr, r, 12, ('=IF(AND(C{0}>0,D{0}=0),"No CRM leads",IF(AND(C{0}>=1500,D{0}>0,F{0}=0),'
                        '"0 sales on "&TEXT(C{0},"$#,##0"),""))').format(r), body_font, fill=fill)
    tier_rules(vr, "A{}:A{}".format(f, l), "$A{}".format(f))
    tier_rules(vr, "K{}:K{}".format(f, l), "$A{}".format(f))

    # ---- Forecast
    title = forecast_title or facts["labels"]["forecast_title"]
    masthead(fc, title, 9, SKY_BLUE)
    put(fc, 3, 1, facts["labels"]["forecast_subtitle"], note_font, bd=False)
    header(fc, 4, ["Month", "2024 Actual", "2025 Actual", "Baseline Forecast", "Planning Low (-12%)",
                   "Planning High (+12%)", "Seasonal prior (DigitalCLIQ estimate)", "+25% Budget", "-15% Budget"])
    for i in range(12):
        r = 5 + i
        fill = alt_fill if (r % 2 == 0) else None
        put(fc, r, 1, "{} 2026".format(MON[i]), label_font, fill=fill)
        put(fc, r, 2, "={}!C{}".format(HN, 5 + i), fmt=NUM, fill=fill)
        put(fc, r, 3, "={}!C{}".format(HN, T1 + 1 + i), fmt=NUM, fill=fill)
        put(fc, r, 4, facts["forecast"]["baseline"][i], fmt=NUM, fill=fill)
        put(fc, r, 5, "=ROUND(D{}*0.88,0)".format(r), fmt=NUM, fill=fill)
        put(fc, r, 6, "=ROUND(D{}*1.12,0)".format(r), fmt=NUM, fill=fill)
        put(fc, r, 7, PRIOR[i], fmt="0.00", fill=fill)
        put(fc, r, 8, "=ROUND(D{0}*(1+{1}!$B$14*{1}!$B$9),0)".format(r, SN), fmt=NUM, fill=fill)
        put(fc, r, 9, "=ROUND(D{0}*(1+{1}!$D$14*{1}!$B$9),0)".format(r, SN), fmt=NUM, fill=fill)
    put(fc, 17, 1, "TOTAL", total_font, fill=total_fill)
    for c in (2, 3, 4, 5, 6, 8, 9):
        col = chr(64 + c)
        put(fc, 17, c, "=SUM({0}5:{0}16)".format(col), total_font, NUM, total_fill)
    put(fc, 17, 7, None, total_font, fill=total_fill)
    put(fc, 18, 1, "vs 2025", label_font)
    for c in (4, 8, 9):
        col = chr(64 + c)
        put(fc, 18, c, "={0}17/C17-1".format(col), fmt=SPCT)
    put(fc, 19, 1, "Planning range, not a statistical interval.", note_font, bd=False)

    # ---- Scenario Planner
    masthead(sp, "Budget Scenario Planner", 5, DIGITAL_BLUE)
    put(sp, 3, 1, "Edit the blue cells. All results are live formulas.", note_font, bd=False)
    put(sp, 4, 1, "MODEL INPUTS", label_font, bd=False)
    lab = el["label"]
    inputs = [
        ("Baseline Spend (2025 actual)", "={}!B{}".format(HN, T2), CUR),
        ("Baseline Leads (2026 forecast)", "={}!D17".format(FN), NUM),
        ("Baseline Sales", "=ROUND(B6*B8,0)", NUM),
        ("Baseline sales per lead (2025 actual; assumed: close rate held at 2025 level)",
         "={0}!G{1}/{0}!C{1}".format(HN, T2), "0.0000"),
        ("Lead Elasticity ({}; p = {:.2f}, n = {})".format(lab, lead_el["p"], lead_el["n"]), lead_el["slope"], "0.000"),
    ]
    for i, (t, v, fm) in enumerate(inputs):
        put(sp, 5 + i, 1, t, label_font)
        c = put(sp, 5 + i, 2, v, input_font, fm, PatternFill("solid", fgColor=CALLOUT_TINT))
        c.border = input_border
    put(sp, 10, 1, "95% CI {:.2f} to {:.2f}. A range that crosses zero means the data cannot rule out no effect.".format(
        lead_el["ci95"][0], lead_el["ci95"][1]), note_font, bd=False)
    put(sp, 12, 1, "SCENARIO RESULTS", label_font, bd=False)
    for i, t in enumerate(["Metric", "+25% Budget", "Status Quo", "-15% Budget", "Custom"], 1):
        put(sp, 13, i, t, hdr_font, fill=band_fill).alignment = ctr
    put(sp, 14, 1, "Budget Adjustment %", label_font)
    for c, v in ((2, 0.25), (3, 0), (4, -0.15), (5, 0.10)):
        cell = put(sp, 14, c, v, input_font if c == 5 else body_font, SPCT,
                   PatternFill("solid", fgColor=CALLOUT_TINT) if c == 5 else None)
    labels = [(15, "Total Budget", "=$B$5*(1+{c}14)", CUR), (16, "Budget Change ($)", "={c}15-$B$5", CUR),
              (17, "Projected Leads", "=ROUND($B$6*(1+{c}14*$B$9),0)", NUM),
              (18, "Lead Change", "={c}17/$B$6-1", SPCT),
              (19, "Projected Sales (assumed: close rate held at 2025 level)", "=ROUND({c}17*$B$8,0)", NUM),
              (20, "Cost per Lead", '=IF({c}17>0,{c}15/{c}17,"-")', CUR2),
              (21, "Cost per Sale", '=IF({c}19>0,{c}15/{c}19,"-")', CUR)]
    for r, t, ftxt, fm in labels:
        put(sp, r, 1, t, label_font)
        for c in "BCDE":
            put(sp, r, "ABCDE".index(c) + 1, ftxt.format(c=c), body_font, fm)
    put(sp, 23, 1, "The elasticity is {}: budget size alone shows no reliable effect on leads. "
        "See the Reallocation Strategy tab.".format(lab.lower()), note_font, bd=False)

    # ---- Executive Summary
    masthead(es, "Test Motors: Lead Forecast and Budget Scenarios", 6, DIGITAL_BLUE)
    put(es, 4, 1, "24-Month Analysis (2024 to 2025) | 2026 Forecast | Budget Scenarios", label_font, bd=False)
    put(es, 5, 1, "Prepared by DigitalCLIQ | Digital Strategy & Development", note_font, bd=False)
    dw = defs["data_window"]
    put(es, 7, 1, "Lead definition: {} {}; excludes {}; data {} to {}.".format(
        defs["report_name"], defs["lead_definition_column"], ", ".join(defs["excluded_sources"]) or "none",
        dw["start"], dw["end"]), body_font, bd=False)
    for i, t in enumerate(["Metric", "2024 Actual", "2025 Actual", "2026 Baseline", "2026 +25% Budget",
                           "2026 -15% Budget"], 1):
        put(es, 9, i, t, hdr_font, fill=band_fill).alignment = ctr
    kp = [(10, "Total Marketing Spend", "B", CUR, "15"), (11, "Total Leads", "C", NUM, "17"),
          (12, "Total Sales", "G", NUM, "19")]
    for r, t, hc, fm, spr in kp:
        put(es, r, 1, t, label_font)
        put(es, r, 2, "={}!{}{}".format(HN, hc, T1), fmt=fm)
        put(es, r, 3, "={}!{}{}".format(HN, hc, T2), fmt=fm)
        put(es, r, 4, "={}!C{}".format(SN, spr), fmt=fm)
        put(es, r, 5, "={}!B{}".format(SN, spr), fmt=fm)
        put(es, r, 6, "={}!D{}".format(SN, spr), fmt=fm)
    put(es, 13, 1, "Close Rate (forecast: held at 2025 level)", label_font)
    put(es, 13, 2, "={}!H{}".format(HN, T1), fmt=PCT)
    put(es, 13, 3, "={}!H{}".format(HN, T2), fmt=PCT)
    for c in (4, 5, 6):
        put(es, 13, c, "=$C$13", fmt=PCT)
    for r, t, ftxt, fm in ((14, "Cost per Lead", '=IF({c}11>0,{c}10/{c}11,"-")', CUR2),
                           (15, "Cost per Sale", '=IF({c}12>0,{c}10/{c}12,"-")', CUR)):
        put(es, r, 1, t, label_font)
        for c in "BCDEF":
            put(es, r, "ABCDEF".index(c) + 1, ftxt.format(c=c), fmt=fm)
    put(es, 16, 1, "YoY Lead Change", label_font)
    put(es, 16, 3, "=C11/B11-1", fmt=SPCT)
    for c in "DEF":
        put(es, 16, "ABCDEF".index(c) + 1, "={}11/C11-1".format(c), fmt=SPCT)
    put(es, 18, 1, "Key finding: spend elasticity {:.3f} ({}; p = {:.2f}, n = {}).".format(
        lead_el["slope"], lab, lead_el["p"], lead_el["n"]), label_font, fill=alt_fill)
    put(es, 19, 1, "Recommendation: move money from zero-sale paid lines to TIER 1 and 2 vendors.", body_font, bd=False)
    put(es, 20, 1, "Forecast: {}. Planning range, not a statistical interval. Tier cut-offs are a DigitalCLIQ "
        "planning assumption. Seasonal prior is a DigitalCLIQ estimate.".format(title.split(": ", 1)[1]),
        note_font, bd=False)

    # ---- Budget Detail
    masthead(bd, "2025 Budget Detail (gross spend)", 15, WARM_GREY)
    header(bd, 4, ["Category", "Vendor"] + MON + ["Annual Total"])
    r = 5
    for name, months in budget:
        fill = alt_fill if (r % 2 == 0) else None
        put(bd, r, 1, "Marketing", label_font, fill=fill)
        put(bd, r, 2, name, label_font, fill=fill)
        for i, m in enumerate(months):
            put(bd, r, 3 + i, m, fmt=CUR, fill=fill)
        put(bd, r, 15, "=SUM(C{0}:N{0})".format(r), fmt=CUR, fill=fill)
        r += 1
    put(bd, r, 1, "TOTAL", total_font, fill=total_fill)
    put(bd, r, 2, None, total_font, fill=total_fill)
    for c in range(3, 16):
        col = chr(64 + c)
        put(bd, r, c, "=SUM({0}5:{0}{1})".format(col, r - 1), total_font, CUR, total_fill)

    # ---- All Providers
    masthead(ap, "All Providers: 24 Months", 9, TILE_BLUE)
    header(ap, 4, ["Lead Provider", "2024 Leads", "2025 Leads", "Total Leads", "Good Leads", "Total Sales",
                   "Close Rate", "YoY Change", "Status"])
    provs = [(s["name"], int(s["leads"] * 0.45), s["leads"] - int(s["leads"] * 0.45), s["leads"], s["sales"],
              "Included" if s["included"] else "Excluded") for s in facts["sources"]]
    provs.sort(key=lambda x: -x[3])
    for i, (n, a, b2, t, sa, st) in enumerate(provs):
        r = 5 + i
        put(ap, r, 1, n, label_font)
        put(ap, r, 2, a, fmt=NUM)
        put(ap, r, 3, b2, fmt=NUM)
        put(ap, r, 4, "=B{0}+C{0}".format(r), fmt=NUM)
        put(ap, r, 5, t, fmt=NUM)
        put(ap, r, 6, sa, fmt=NUM)
        put(ap, r, 7, '=IF(E{0}>0,F{0}/E{0},"-")'.format(r), fmt=PCT)
        put(ap, r, 8, '=IF(B{0}>0,C{0}/B{0}-1,"-")'.format(r), fmt=SPCT)
        put(ap, r, 9, st, body_font)

    # ---- Reallocation
    masthead(rs, "Reallocation Strategy: Budget Neutral", 7, TILE_BLUE)
    put(rs, 3, 1, "Section 2: proposed reallocation (CUT order from verify_workbook.py tiers)", label_font, bd=False)
    header(rs, 4, ["Action", "Vendor", "Tier", "Current", "Proposed", "Change", "Rationale"])
    cut = [v["vendor"] for v in vw.cut_order(rows)]
    plan = []
    moved = 0
    for v in rows:
        if v["vendor"] in cut:
            new = int(v["spend"] * 0.5)
            moved += v["spend"] - new
            plan.append(("CUT", v, new, "Cut: " + (v["note"] or v["tier"])))
    invest = [v for v in rows if v["tier"] in ("TIER 1", "TIER 2")]
    for j, v in enumerate(invest):
        add = moved // len(invest) + (moved % len(invest) if j == 0 else 0)
        plan.append(("INVEST", v, v["spend"] + add, "Invest: " + v["tier"]))
    order = {name: i for i, name in enumerate(cut)}
    plan.sort(key=lambda p: (p[0] != "CUT", order.get(p[1]["vendor"], 99)))
    r = 5
    for act, v, new, why in plan:
        spend_fill = (PatternFill("solid", fgColor=WARM_GREY) if act == "CUT" else PatternFill("solid", fgColor=CALLOUT_TINT))
        spend_font = Font(name=B, bold=True, size=10, color=WHITE if act == "CUT" else DIGITAL_BLUE)
        put(rs, r, 1, act, label_font)
        put(rs, r, 2, v["vendor"], label_font)
        put(rs, r, 3, v["tier"], body_font)
        put(rs, r, 4, v["spend"], spend_font, CUR, spend_fill)
        put(rs, r, 5, new, spend_font, CUR, spend_fill)
        put(rs, r, 6, "=E{0}-D{0}".format(r), fmt=CUR)
        put(rs, r, 7, why, body_font)
        r += 1
    put(rs, r, 1, "TOTAL", total_font, fill=total_fill)
    for c in (4, 5, 6):
        col = chr(64 + c)
        put(rs, r, c, "=SUM({0}5:{0}{1})".format(col, r - 1), total_font, CUR, total_fill)
    put(rs, r + 2, 1, "Methodology: tier cut-offs are a DigitalCLIQ planning assumption; projections use each "
        "vendor's own CPL and close rate.", note_font, bd=False)

    wb.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1]
    ran = sys.argv[2].split(",") if len(sys.argv) > 2 else ["prophet", "holt_winters", "seasonal_prior"]
    title = sys.argv[3] if len(sys.argv) > 3 else None
    facts, rows, budget = make_data(models_ran=ran)
    build(facts, rows, budget, out, title)
    with open(out[:-5] + ".facts.json", "w") as fh:
        json.dump(facts, fh, indent=1)
    print("built", out)
